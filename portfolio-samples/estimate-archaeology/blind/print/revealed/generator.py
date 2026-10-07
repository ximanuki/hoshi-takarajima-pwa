#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SEALED -- generator for the "print" blind dataset (small Japanese commercial print shop).

    python3 generator.py

Env overrides (used only by the self-check): EA_PUBLIC_DIR, EA_SEALED_DIR.
Deterministic: fixed seed, no wall-clock or hash-order dependence; xlsx zips are normalised.

Public outputs  (EA_PUBLIC_DIR): quotes.csv, SCHEMA.md, messy/*.xlsx, SEALED.sha256
Sealed outputs  (EA_SEALED_DIR): RULES.md, truth.csv   (+ this file)
"""
import csv
import datetime as dt
import hashlib
import io
import math
import os
import re
import zipfile
from decimal import Decimal, ROUND_HALF_UP
from fractions import Fraction as F

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.pagebreak import Break

SEED = 7310429
N_BASE = 2373          # + duplicates + re-issues = 2,400 rows

HERE = os.path.dirname(os.path.abspath(__file__))
PUBLIC_DIR = os.environ.get(
    "EA_PUBLIC_DIR",
    "/home/user/hoshi-takarajima-pwa/portfolio-samples/estimate-archaeology/blind/print")
SEALED_DIR = os.environ.get("EA_SEALED_DIR", HERE)

D = dt.date
START, END = D(2022, 10, 1), D(2025, 9, 30)
REV_DATES = [D(2023, 4, 1), D(2024, 7, 1), D(2025, 4, 1)]
INVOICE_DATE = D(2023, 10, 1)
WATANABE_START = D(2023, 6, 1)
SUZUKI_END = D(2025, 3, 31)
FIXED_TS = (2025, 10, 1, 9, 0, 0)


def period_of(d):
    return sum(1 for r in REV_DATES if d >= r)


# ----------------------------------------------------------------------------- calendar
HOLIDAYS = {D.fromisoformat(s) for s in """
2022-10-10 2022-11-03 2022-11-23
2023-01-09 2023-02-23 2023-03-21 2023-05-03 2023-05-04 2023-05-05 2023-07-17 2023-08-11
2023-09-18 2023-10-09 2023-11-03 2023-11-23
2024-01-08 2024-02-12 2024-02-23 2024-03-20 2024-04-29 2024-05-03 2024-05-06 2024-07-15
2024-08-12 2024-09-16 2024-09-23 2024-10-14 2024-11-04
2025-01-13 2025-02-11 2025-02-24 2025-03-20 2025-04-29 2025-05-05 2025-05-06 2025-07-21
2025-08-11 2025-09-15 2025-09-23
""".split()}


def is_open(d):
    if d.weekday() == 6 or d in HOLIDAYS:
        return False
    if (d.month == 12 and d.day >= 29) or (d.month == 1 and d.day <= 4):
        return False
    if d.month == 8 and 13 <= d.day <= 16:
        return False
    return True


def is_bday(d):
    return is_open(d) and d.weekday() < 5


def add_bdays(d, n):
    while n > 0:
        d += dt.timedelta(1)
        if is_bday(d):
            n -= 1
    return d


def ceil_div(a, b):
    return -(-a // b)


def floor_to(x, unit):
    return math.floor(F(x) / unit) * unit


def ceil_to(x, unit):
    return math.ceil(F(x) / unit) * unit


def half_up_to(x, unit):
    return math.floor(F(x) / unit + F(1, 2)) * unit


def dq(x, q):
    return Decimal(str(x)).quantize(Decimal(q), rounding=ROUND_HALF_UP)


def pick(rng, items, p=None):
    if p is None:
        return items[int(rng.integers(len(items)))]
    p = np.asarray(p, dtype=float)
    return items[int(rng.choice(len(items), p=p / p.sum()))]


# ----------------------------------------------------------------------------- price tables
PAPER_KIKU_P0 = {"コート90kg": "12.6", "コート110kg": "15.2", "マットコート135kg": "19.8",
                 "上質70kg": "10.4", "上質90kg": "13.1"}
PAPER_STEP = [None, "1.15", "1.08", "1.00"]
ENV_BOX_P0 = {("長3", "クラフト70g"): 2400, ("長3", "白ケント80g"): 4000,
              ("洋長3", "白ケント80g"): 5600, ("角2", "クラフト85g"): 4800,
              ("角2", "白ケント100g"): 7000}
ENV_STEP = [None, "1.12", "1.06", "1.05"]
ENV_PACK = {"長3": 1000, "洋長3": 1000, "角2": 500}
CARD_FIRST_P0 = {"1/0": 1600, "4/0": 2400, "1/1": 2600, "4/4": 3400}
CARD_ADD_P0 = {"1/0": 400, "4/0": 700, "1/1": 800, "4/4": 1000}
SPECIAL_CARD_PAPER = ("ヴァンヌーボ195kg", "パール紙180kg")


def build_tables():
    kiku = [dict(PAPER_KIKU_P0)]
    env = [dict(ENV_BOX_P0)]
    for p in range(1, 4):
        kiku.append({k: str(dq(Decimal(v) * Decimal(PAPER_STEP[p]), "0.1")) for k, v in kiku[-1].items()})
        env.append({k: int(dq(Decimal(v) * Decimal(ENV_STEP[p]), "1E1")) for k, v in env[-1].items()})
    out = []
    for p in range(4):
        old = p < 2
        T = {
            "plate": 2500 if old else 2800,
            "imp_base": 2000 if old else 2200,
            "imp_add": 700 if old else 750,
            "cut": 800 if old else 1000,
            "click_color": 30 if old else 32,
            "click_mono": 8 if old else 9,
            "data_fee": 1500,
            "min": 5000 if old else 5500,
            "round_unit": 10 if old else 100,
            "fix_fee": 3000 if old else 3500,
            "fix_fee_card": 1000,
            "spot_fee": 4000,
            "proof_off": 8000,
            "proof_od": 1500,
            "band_fee": 400,
            "ship_city": 800, "ship_city_free": 10000,
            "ship_out": 1500, "ship_out_free": 30000,
            "ship_box": 1300 if old else 1500,
            "ship_mail": 520 if old else 600,
            "bind_setup": 3000 if old else 3500,
            "bind_small": 10 if old else 12,
            "bind_large": 15 if old else 18,
            "pp_setup": 3000,
            "pp_a4": 15 if old else 16,
            "fold": {"二つ折": (2000, 1200), "三つ折": (2500, 1600), "Z折": (2500, 1600)},
            "env_plate": 1800 if old else 2000,
            "env_imp_base": 2000 if old else 2200,
            "env_imp_add": 600 if old else 650,
            "card_first": {k: v + (200 if p >= 3 else 0) for k, v in CARD_FIRST_P0.items()},
            "card_add": {k: v + (100 if p >= 3 else 0) for k, v in CARD_ADD_P0.items()},
            "paper_kiku": {k: F(v) for k, v in kiku[p].items()},
            "paper_46": {k: F(str(dq(Decimal(v) * Decimal("1.38"), "0.1"))) for k, v in kiku[p].items()},
            "paper_od": {k: F(str(dq(Decimal(v) * Decimal("0.45"), "0.1"))) for k, v in kiku[p].items()},
            "env_box": env[p],
        }
        out.append(T)
    return out


TABLES = build_tables()

UP_OFF = {"A3": 2, "A4": 4, "A5": 8, "B4": 4, "B5": 8}
UP_OD = {"A3": 1, "A4": 2, "A5": 4, "B4": 1, "B5": 2}
A4EQ = {"A3": F(2), "A4": F(1), "A5": F(1, 2), "B4": F(3, 2), "B5": F(3, 4)}
GSM = {"コート90kg": "104.7", "コート110kg": "128", "マットコート135kg": "157", "上質70kg": "81.4",
       "上質90kg": "104.7", "マットコート220kg": "256", "上質180kg": "209.3", "上質135kg": "157",
       "ヴァンヌーボ195kg": "227", "パール紙180kg": "209.3"}
AREA = {"A3": "0.12474", "A4": "0.06237", "A5": "0.031185", "B4": "0.093548", "B5": "0.046774",
        "91x55": "0.005005"}
ENV_G = {("長3", "クラフト70g"): "4.5", ("長3", "白ケント80g"): "5.2", ("洋長3", "白ケント80g"): "5.6",
         ("角2", "クラフト85g"): "14", ("角2", "白ケント100g"): "16.5"}

RUSH_RE = re.compile("特急")
FIX_RE = re.compile("データ修正|文字修正|修正あり")
SPOT_RE = re.compile("特色|DIC")
PROOF_RE = re.compile("色校")
BAND_RE = re.compile("帯掛")
REPEAT_RE = re.compile("前回同様|リピート")
DESIGN_RE = re.compile("新規デザイン|新規作成")

COMP_KEYS = ["plates", "impressions", "paper", "clicks", "data_fee", "cutting", "folding", "pp",
             "binding", "envelopes", "cards"]
EXTRA_KEYS = ["data_fix", "spot", "proof", "banding", "design"]


def parse_colors(c):
    a, b = c.split("/")
    return int(a), int(b)


def method_of(spec):
    p = spec["product"]
    if p == "名刺":
        return "card"
    if p == "封筒":
        return "env"
    if p == "冊子":
        return "od" if spec["qty"] < 300 else "offset"
    return "od" if spec["qty"] < 1000 else "offset"


def click_units(sides):
    a = min(sides, 300)
    b = min(max(sides - 300, 0), 700)
    c = max(sides - 1000, 0)
    return F(a) + F(b) * F(4, 5) + F(c) * F(3, 5)


def weight_g(spec):
    prod, qty, kinds = spec["product"], spec["qty"], spec["kinds"]
    if prod == "封筒":
        return F(ENV_G[(spec["size"], spec["paper"])]) * qty * kinds
    if prod == "冊子":
        return F(qty * (spec["pages"] // 2)) * F(AREA[spec["size"]]) * F(GSM[spec["paper"]])
    return F(qty * kinds) * F(AREA[spec["size"]]) * F(GSM[spec["paper"]])


def compute(spec, notes, date, rep, rate, *, period=None, force_rush=False, skip_delivery=False,
            rounder=None):
    """The shop's estimating logic. rate: str/Fraction customer multiplier."""
    p = period_of(date) if period is None else period
    T = TABLES[p]
    r = F(rate)
    prod, qty, kinds = spec["product"], spec["qty"], spec["kinds"]
    pieces = qty * kinds
    m = method_of(spec)
    c = {k: F(0) for k in COMP_KEYS}
    if prod in ("チラシ", "リーフレット"):
        size = spec["size"]
        cf, cb = parse_colors(spec["colors"])
        if m == "offset":
            ppk = cf + cb
            ts = ceil_div(qty, UP_OFF[size])
            c["plates"] = F(T["plate"] * ppk * kinds)
            per_plate = T["imp_base"] + T["imp_add"] * (ceil_div(ts, 1000) - 1)
            c["impressions"] = F(per_plate * ppk * kinds)
            press = kinds * (ts + 30 * ppk + ceil_div(ts * 2, 100))
            billed = ceil_div(ceil_div(press, 2), 250) * 250
            price = T["paper_kiku" if size[0] == "A" else "paper_46"][spec["paper"]]
            c["paper"] = billed * price
        else:
            S = kinds * ceil_div(qty, UP_OD[size])
            spoil = max(3, ceil_div(S * 3, 100))
            c["paper"] = (S + spoil) * T["paper_od"][spec["paper"]]
            sides = S * (2 if cb else 1)
            c["clicks"] = click_units(sides) * (T["click_color"] if cf == 4 else T["click_mono"])
            c["data_fee"] = F(0 if REPEAT_RE.search(notes) else T["data_fee"] * kinds)
        c["cutting"] = F(T["cut"])
        fin = spec["finishing"]
        if fin in T["fold"]:
            s0, per = T["fold"][fin]
            c["folding"] = F(s0 + per * ceil_div(pieces, 1000))
        if "PP" in fin:
            c["pp"] = T["pp_setup"] + T["pp_a4"] * A4EQ[size] * pieces
    elif prod == "冊子":
        size, pages, paper = spec["size"], spec["pages"], spec["paper"]
        cf, cb = parse_colors(spec["colors"])
        if m == "od":
            spc = pages // 4 if size == "A4" else pages // 8
            S = qty * spc
            spoil = max(3, ceil_div(S * 3, 100))
            c["paper"] = (S + spoil) * T["paper_od"][paper]
            c["clicks"] = click_units(2 * S) * (T["click_color"] if cf == 4 else T["click_mono"])
            c["data_fee"] = F(0 if REPEAT_RE.search(notes) else T["data_fee"])
        else:
            per_full = 8 if size == "A4" else 16
            full, half = pages // per_full, (pages % per_full) // (per_full // 2)
            plates = full * (cf + cb) + half * cf
            c["plates"] = F(T["plate"] * plates)
            per_plate = T["imp_base"] + T["imp_add"] * (ceil_div(qty, 1000) - 1)
            c["impressions"] = F(per_plate * plates)
            net_sheets = full * qty + half * ceil_div(qty, 2)
            press = net_sheets + 30 * plates + ceil_div(net_sheets * 2, 100)
            billed = ceil_div(ceil_div(press, 2), 250) * 250
            c["paper"] = billed * T["paper_kiku"][paper]
        per_copy = T["bind_small"] if pages <= 16 else T["bind_large"]
        c["binding"] = F(T["bind_setup"] + per_copy * qty)
    elif m == "env":
        size, paper = spec["size"], spec["paper"]
        ncol = parse_colors(spec["colors"])[0]
        c["envelopes"] = F(ceil_div(pieces, ENV_PACK[size]) * T["env_box"][(size, paper)])
        c["plates"] = F(T["env_plate"] * ncol * kinds)
        per = T["env_imp_base"] + T["env_imp_add"] * (ceil_div(qty, 1000) - 1)
        c["impressions"] = F(per * ncol * kinds) * (F(3, 2) if size == "角2" else F(1))
    else:  # card
        boxes = qty // 100
        fin = spec["finishing"]
        per = T["card_first"][spec["colors"]] + T["card_add"][spec["colors"]] * (boxes - 1)
        if spec["paper"] in SPECIAL_CARD_PAPER:
            per += 300 * boxes
        if "PP" in fin:
            per += 600 * boxes
        if "角丸" in fin:
            per += 500
        tot = F(0)
        for i in range(1, kinds + 1):
            tot += per * (F(1) if i <= 5 else (F(9, 10) if i <= 10 else F(4, 5)))
        c["cards"] = tot

    base = sum(c.values(), F(0))
    after_rate = base * r
    min_applied = (m != "card") and after_rate < T["min"]
    after_min = F(T["min"]) if min_applied else after_rate
    rush = after_min / 4 if (force_rush or RUSH_RE.search(notes)) else F(0)
    ex = {k: F(0) for k in EXTRA_KEYS}
    if FIX_RE.search(notes):
        ex["data_fix"] = F(T["fix_fee_card"] if m == "card" else T["fix_fee"])
    if SPOT_RE.search(notes) and m in ("offset", "env"):
        ex["spot"] = F(T["spot_fee"])
    if PROOF_RE.search(notes) and m in ("offset", "od"):
        ex["proof"] = F(T["proof_off"] if m == "offset" else T["proof_od"])
    if BAND_RE.search(notes) and m != "card":
        ex["banding"] = F(T["band_fee"] * ceil_div(pieces, 1000))
    if m == "card" and DESIGN_RE.search(notes):
        ex["design"] = F(3000 + 500 * (kinds - 1))
    pre_round = after_min + rush + sum(ex.values(), F(0))
    unit = T["round_unit"]
    sub = rounder(pre_round, unit) if rounder else floor_to(pre_round, unit)

    dlv = spec["delivery"]
    if dlv == "引取":
        fee = 0
    elif dlv == "市内":
        fee = 0 if sub >= T["ship_city_free"] else T["ship_city"]
    elif dlv == "市外":
        fee = 0 if sub >= T["ship_out_free"] else T["ship_out"]
    else:
        kg = weight_g(spec) / 1000
        fee = T["ship_mail"] if kg <= 1 else T["ship_box"] * math.ceil(kg / 15)
    would_fee = fee
    if skip_delivery:
        fee = 0
    net0 = sub + fee
    net = (net0 // 500) * 500 if (rep == "高橋" and net0 >= 10000) else net0
    return {"method": m, "period": p, "rate": r, "comp": c, "base": base, "after_rate": after_rate,
            "min_applied": min_applied, "after_min": after_min, "rush": rush, "extras": ex,
            "pre_round": pre_round, "unit": unit, "sub": int(sub), "delivery": int(fee),
            "would_delivery": int(would_fee), "net0": int(net0), "rep_adj": int(net - net0),
            "net": int(net)}


# ----------------------------------------------------------------------------- customers / specs
PRODUCTS = ["チラシ", "リーフレット", "冊子", "名刺", "封筒"]
INDUSTRY_MIX = {
    "realestate": [.80, .05, .02, .10, .03],
    "retail": [.55, .15, .05, .20, .05],
    "clinic": [.15, .25, .05, .30, .25],
    "office": [.10, .10, .10, .45, .25],
    "maker": [.15, .20, .15, .30, .20],
    "school": [.30, .05, .40, .05, .20],
    "agency": [.45, .25, .20, .05, .05],
    "individual": [.30, .05, .15, .50, .00],
    "walkin": [.30, .05, .05, .60, .00],
}
QTY_LADDER = {
    "flat": [50, 100, 200, 300, 500, 1000, 1500, 2000, 3000, 5000, 8000, 10000, 15000, 20000, 30000, 50000],
    "冊子": [30, 50, 100, 150, 200, 300, 500, 1000, 1500, 2000, 3000],
    "名刺": [100, 200, 300, 500, 1000],
    "封筒": [300, 500, 1000, 1200, 1500, 2000, 2500, 3000, 5000, 10000],
}


def ladder_key(prod):
    return "flat" if prod in ("チラシ", "リーフレット") else prod


def make_spec(rng, prod, ind):
    s = {"product": prod, "size": "", "paper": "", "colors": "", "pages": 0, "finishing": "",
         "qty": 0, "kinds": 1}
    if prod == "チラシ":
        sz = {"realestate": (["B4", "A4", "A3", "B5"], [.35, .45, .1, .1]),
              "retail": (["A4", "B5", "A5", "A3", "B4"], [.5, .1, .15, .1, .15]),
              "school": (["A4", "B5", "A5", "B4"], [.6, .2, .1, .1])}.get(
            ind, (["A4", "A5", "A3", "B5"], [.65, .2, .05, .1]))
        s["size"] = pick(rng, *sz)
        pp = {"realestate": (["コート90kg", "コート110kg", "上質70kg"], [.7, .2, .1]),
              "retail": (["コート90kg", "コート110kg", "マットコート135kg", "上質70kg"], [.45, .25, .2, .1]),
              "school": (["上質70kg", "上質90kg", "コート90kg"], [.55, .2, .25]),
              "individual": (["コート110kg", "マットコート135kg", "上質90kg"], [.4, .3, .3]),
              "walkin": (["コート110kg", "マットコート135kg", "上質90kg"], [.4, .3, .3])}.get(
            ind, (["コート90kg", "コート110kg", "マットコート135kg", "上質90kg"], [.4, .3, .2, .1]))
        s["paper"] = pick(rng, *pp)
        cc = {"realestate": (["4/4", "4/0"], [.55, .45]),
              "school": (["1/0", "1/1", "4/0"], [.45, .25, .3])}.get(
            ind, (["4/0", "4/4", "1/0", "1/1"], [.45, .35, .12, .08]))
        s["colors"] = pick(rng, *cc)
        qq = {"realestate": ([3000, 5000, 10000, 20000, 30000, 50000], [.1, .25, .3, .2, .1, .05]),
              "retail": ([300, 500, 1000, 2000, 3000, 5000, 10000], [.1, .15, .25, .2, .12, .12, .06]),
              "school": ([100, 200, 300, 500, 1000, 1500, 2000], [.1, .2, .2, .25, .15, .05, .05]),
              "individual": ([50, 100, 200, 300, 500, 1000], [.1, .3, .25, .15, .15, .05]),
              "walkin": ([50, 100, 200, 300, 500, 1000], [.1, .3, .25, .15, .15, .05]),
              "agency": ([1000, 2000, 3000, 5000, 10000, 20000, 30000], [.1, .15, .2, .25, .2, .07, .03])}.get(
            ind, ([200, 300, 500, 1000, 2000, 3000, 5000], [.08, .12, .2, .25, .15, .1, .1]))
        s["qty"] = pick(rng, *qq)
        u = rng.random()
        if s["paper"] == "マットコート135kg" and s["qty"] <= 1000 and u < 0.12:
            s["finishing"] = "片面PP"
        elif u > 0.98:
            s["finishing"] = "二つ折"
        s["kinds"] = pick(rng, [1, 2, 3], [.8, .12, .08] if ind == "realestate" else [.94, .05, .01])
    elif prod == "リーフレット":
        s["size"] = pick(rng, ["A4", "A3", "B4"], [.75, .2, .05])
        if s["size"] == "A4":
            s["finishing"] = pick(rng, ["三つ折", "Z折", "二つ折"], [.65, .2, .15])
        else:
            s["finishing"] = pick(rng, ["二つ折", "三つ折"], [.85, .15])
        s["paper"] = pick(rng, ["コート110kg", "マットコート135kg", "コート90kg"], [.5, .35, .15])
        s["colors"] = pick(rng, ["4/4", "4/0", "1/1"], [.9, .05, .05])
        s["qty"] = pick(rng, [300, 500, 1000, 2000, 3000, 5000, 10000], [.08, .15, .3, .2, .12, .1, .05])
        s["kinds"] = pick(rng, [1, 2], [.97, .03])
    elif prod == "冊子":
        s["size"] = pick(rng, ["A4", "A5"], [.7, .3])
        if s["size"] == "A4":
            s["pages"] = pick(rng, [8, 12, 16, 20, 24, 28, 32], [.2, .2, .25, .12, .12, .05, .06])
        else:
            s["pages"] = pick(rng, [8, 16, 24, 32], [.35, .35, .2, .1])
        if ind == "school":
            s["paper"] = pick(rng, ["上質70kg", "コート110kg", "マットコート135kg"], [.5, .3, .2])
            s["colors"] = pick(rng, ["1/1", "4/4"], [.45, .55])
            s["qty"] = pick(rng, [50, 100, 150, 200, 300, 500], [.1, .25, .15, .2, .15, .15])
        elif ind in ("individual", "walkin"):
            s["paper"] = pick(rng, ["上質70kg", "コート110kg", "マットコート135kg"], [.3, .4, .3])
            s["colors"] = pick(rng, ["1/1", "4/4"], [.3, .7])
            s["qty"] = pick(rng, [30, 50, 100, 200], [.3, .35, .25, .1])
        else:
            s["paper"] = pick(rng, ["コート110kg", "マットコート135kg", "上質70kg"], [.45, .4, .15])
            s["colors"] = pick(rng, ["4/4", "1/1"], [.9, .1])
            s["qty"] = pick(rng, [100, 200, 300, 500, 1000, 2000, 3000], [.12, .15, .2, .25, .15, .08, .05])
        s["finishing"] = "中綴じ"
    elif prod == "名刺":
        s["size"] = "91x55"
        s["paper"] = pick(rng, ["マットコート220kg", "上質180kg", "ヴァンヌーボ195kg", "パール紙180kg", "上質135kg"],
                          [.42, .30, .15, .05, .08])
        s["colors"] = pick(rng, ["4/0", "4/4", "1/0", "1/1"], [.35, .25, .25, .15])
        s["qty"] = pick(rng, [100, 200, 300, 500, 1000], [.62, .2, .06, .1, .02])
        s["finishing"] = pick(rng, ["", "片面PP", "角丸", "片面PP+角丸"], [.86, .07, .05, .02])
        if ind in ("office", "maker"):
            s["kinds"] = pick(rng, [1, 2, 3, 4, 5, 6, 8, 10, 12, 15, 20],
                              [.25, .15, .12, .1, .08, .07, .07, .06, .04, .04, .02])
        elif ind in ("individual", "walkin"):
            s["kinds"] = pick(rng, [1, 2], [.88, .12])
        else:
            s["kinds"] = pick(rng, [1, 2, 3, 4, 5, 6], [.4, .25, .15, .1, .05, .05])
        if s["kinds"] > 4 and s["qty"] > 200:
            s["qty"] = pick(rng, [100, 200], [.7, .3])
    else:  # 封筒
        s["size"] = pick(rng, ["長3", "角2", "洋長3"], [.55, .35, .10])
        if s["size"] == "長3":
            s["paper"] = pick(rng, ["クラフト70g", "白ケント80g"], [.6, .4])
        elif s["size"] == "角2":
            s["paper"] = pick(rng, ["クラフト85g", "白ケント100g"], [.55, .45])
        else:
            s["paper"] = "白ケント80g"
        s["colors"] = pick(rng, ["1/0", "2/0"], [.65, .35])
        if s["size"] == "角2":
            s["qty"] = pick(rng, [300, 500, 1000, 1200, 1500, 2000, 3000], [.08, .3, .3, .04, .08, .12, .08])
        else:
            s["qty"] = pick(rng, [500, 1000, 1500, 2000, 3000, 5000, 10000], [.1, .35, .05, .2, .12, .13, .05])
        s["kinds"] = pick(rng, [1, 2], [.95, .05])
    return s


def build_customers(rng):
    custs = []
    others = ["realestate"] * 5 + ["retail"] * 10 + ["clinic"] * 7 + ["office"] * 10 + ["maker"] * 6
    perm = rng.permutation(len(others))
    others = [others[i] for i in perm]
    c_inds = ["realestate", others[0], "maker"] + others[1:]
    fac = {"realestate": 1.8, "retail": 1.0, "clinic": 0.7, "office": 0.8, "maker": 1.0, "agency": 1.5,
           "school": 0.8, "individual": 0.2}
    for i, ind in enumerate(c_inds, 1):
        code = "C%03d" % i
        rate = pick(rng, ["1.00", "0.95", "0.90", "0.92", "1.05", "1.10"], [.5, .2, .12, .05, .08, .05])
        rates = [(START, rate)]
        w = float(rng.lognormal(0, 0.7)) * fac[ind]
        if code == "C001":
            rates, w = [(START, "0.90")], 5.0
        if code == "C003":
            rates, w = [(START, "1.00"), (D(2024, 1, 1), "0.95")], max(w, 1.0) * 1.5
        custs.append({"code": code, "ctype": "C", "industry": ind, "rates": rates, "w": w})
    for i, rate in enumerate(["0.80", "0.80", "0.85", "0.75", "0.85", "0.80"], 1):
        code = "D%02d" % i
        rates = [(START, rate)]
        if code == "D02":
            rates = [(START, "0.80"), (D(2024, 7, 1), "0.85")]
        custs.append({"code": code, "ctype": "D", "industry": "agency", "rates": rates,
                      "w": float(rng.lognormal(0, 0.6)) * fac["agency"] * (1.4 if code == "D02" else 1)})
    for i in range(1, 9):
        code = "G%02d" % i
        rate = {5: "1.00", 7: "0.95"}.get(i, "0.90")
        custs.append({"code": code, "ctype": "G", "industry": "school", "rates": [(START, rate)],
                      "w": float(rng.lognormal(0, 0.6)) * fac["school"]})
    for i in range(1, 16):
        custs.append({"code": "K%03d" % i, "ctype": "K", "industry": "individual", "rates": [(START, "1.00")],
                      "w": float(rng.lognormal(0, 0.5)) * fac["individual"]})
    tot = sum(c["w"] for c in custs)
    custs.append({"code": "一般", "ctype": "W", "industry": "walkin", "rates": [(START, "1.00")],
                  "w": tot * 0.11 / 0.89})
    always = ("C001", "C003", "D02", "G01", "一般")
    for c in custs:
        c["start"], c["end"] = START, END
        if c["code"] not in always:
            if rng.random() < 0.18:
                c["start"] = D(2022, 11, 1) + dt.timedelta(int(rng.integers(0, 880)))
            if rng.random() < 0.10:
                e = D(2023, 6, 1) + dt.timedelta(int(rng.integers(0, 820)))
                if (e - c["start"]).days > 180:
                    c["end"] = e
    # balance account load across reps: heaviest accounts first, mostly to the rep furthest below target
    target = {"田中": 0.30, "鈴木": 0.35, "高橋": 0.35}
    load = {k: 0.0 for k in target}
    order = sorted([c for c in custs if c["ctype"] in "CDG"], key=lambda c: (-c["w"], c["code"]))
    for c in order:
        if rng.random() < 0.3:
            c["r0"] = pick(rng, list(target))
        else:
            tot = sum(load.values()) + c["w"]
            c["r0"] = max(target, key=lambda k: (target[k] * tot - load[k], k))
        load[c["r0"]] += c["w"]
    for c in custs:
        if c["ctype"] in "CDG":
            r0 = c.pop("r0")
        elif c["ctype"] == "K":
            r0 = pick(rng, ["田中", "鈴木"], [.5, .5])
        else:
            r0 = None
        hist = [(START, r0)]
        if r0 == "田中" and c["ctype"] in "CG" and rng.random() < 0.4:
            hist.append((WATANABE_START, "渡辺"))
        u = rng.random()
        if hist[-1][1] == "鈴木":
            hist.append((SUZUKI_END + dt.timedelta(1), "渡辺" if u < 0.65 else "高橋"))
        c["reps"] = hist
        ind = c["industry"]
        if ind in ("maker", "agency"):
            c["home"] = pick(rng, ["市内", "市外", "県外"], [.4, .35, .25])
        elif ind == "school":
            c["home"] = pick(rng, ["市内", "市外"], [.7, .3])
        elif ind in ("individual", "walkin"):
            c["home"] = None
        else:
            c["home"] = pick(rng, ["市内", "市外", "県外"], [.6, .3, .1])
        if c["code"] == "C001":
            c["home"] = "市内"
        c["pickup_p"] = {"C": .12, "D": .08, "G": .2}.get(c["ctype"], 0)
        c["templates"] = []
    return custs


def cust_rate(cust, d):
    r = cust["rates"][0][1]
    for since, val in cust["rates"]:
        if d >= since:
            r = val
    return r


def active_reps(d):
    out = ["田中", "高橋"]
    if d <= SUZUKI_END:
        out.append("鈴木")
    if d >= WATANABE_START:
        out.append("渡辺")
    return out


def rep_hist_at(cust, d):
    rep = cust["reps"][0][1]
    for since, val in cust["reps"]:
        if d >= since:
            rep = val
    return rep


def rep_for(rng, cust, d):
    if cust["ctype"] == "W":
        if d < WATANABE_START:
            rep = pick(rng, ["田中", "鈴木"], [.6, .4])
        elif d <= SUZUKI_END:
            rep = pick(rng, ["田中", "鈴木", "渡辺"], [.4, .3, .3])
        else:
            rep = pick(rng, ["田中", "渡辺"], [.5, .5])
    else:
        rep = cust["reps"][0][1]
        for since, val in cust["reps"]:
            if d >= since:
                rep = val
    if rng.random() < 0.04:
        alt = [r for r in active_reps(d) if r != rep]
        rep = pick(rng, alt)
    return rep


# ----------------------------------------------------------------------------- notes
RUSH_PH = ["特急", "特急対応", "特急でお願いします", "特急（翌日納品希望）"]
HURRY_PH = ["至急", "急ぎ", "お急ぎ"]
FIX_PH = ["データ修正あり", "文字修正あり", "文字修正3箇所", "データ修正（電話番号変更）", "修正あり（PDF校正2回）"]
SPOT_PH = ["特色1色（DIC 156）", "ロゴ特色指定", "DIC指定あり", "特色（DIC 2585）"]
ENV_SPOT_PH = ["社名DIC 182", "ロゴDIC指定", "特色（DIC 221）", "2色目DIC指定"]
PROOF_OFF_PH = ["本機色校正あり", "色校1回"]
PROOF_OD_PH = ["色校正出力1部", "色校あり"]
BAND_PH = ["100枚ごと帯掛け", "帯掛け希望", "500枚ごと帯掛け"]
REPEAT_PH = ["前回同様", "リピート", "前回同様（数量のみ変更）"]
DESIGN_PH = ["新規デザイン", "新規作成（ロゴ支給）", "新規デザイン（2案）"]
COMPETE_PH = ["相見積", "他社相見積あり"]
NOISE = ["PDF入稿", "完全データ入稿", "Illustrator入稿", "Word入稿（要変換）", "納品書同梱", "請求書は本社宛",
         "見積のみ", "検討中", "要確認", "サンプル5部別送", "メールにて見積送付", "FAX送付済", "見積有効期限1ヶ月",
         "紙は前回と同じ", "支払は月末締め翌月末", "担当者様不在のため後日連絡", "修正なし", "校正はPDFで"]
NOISE_BY_PROD = {"リーフレット": ["折り見本あり"], "冊子": ["表紙・本文共紙", "ページ数変更の可能性あり"],
                 "名刺": ["裏面英語表記", "肩書変更あり", "QRコード入り"], "封筒": ["〒枠なし", "窓なし", "社名・住所入り"],
                 "チラシ": ["新聞折込用", "ポスティング用", "地図差替えあり"]}


def make_notes(rng, spec, is_repeat, qty_up):
    m = method_of(spec)
    prod = spec["product"]
    ph = []
    rush = hurry = False
    if is_repeat and rng.random() < 0.45:
        ph.append(pick(rng, REPEAT_PH))
    if is_repeat and qty_up and rng.random() < 0.3:
        ph.append("前回より部数増")
    u = rng.random()
    if u < 0.04:
        ph.append(pick(rng, RUSH_PH))
        rush = True
    elif u < 0.07:
        ph.append(pick(rng, HURRY_PH))
        hurry = True
    u = rng.random()
    if u < (0.08 if prod == "名刺" else 0.07):
        ph.append(pick(rng, FIX_PH))
    elif u < 0.12:
        ph.append("修正なし")
    u = rng.random()
    if m == "offset" and u < 0.04:
        ph.append(pick(rng, SPOT_PH))
    elif m == "env" and u < 0.30:
        ph.append(pick(rng, ENV_SPOT_PH if spec["colors"] == "2/0" else ENV_SPOT_PH[:3]))
    elif m == "od" and u < 0.015:
        ph.append("DIC近似色で")
    u = rng.random()
    if m == "offset" and u < 0.025:
        ph.append(pick(rng, PROOF_OFF_PH))
    elif m == "od" and u < 0.012:
        ph.append(pick(rng, PROOF_OD_PH))
    if m in ("offset", "od", "env") and prod != "冊子" and spec["qty"] * spec["kinds"] >= 1000 \
            and rng.random() < 0.03:
        ph.append(pick(rng, BAND_PH))
    if m == "card" and not is_repeat and rng.random() < 0.25:
        ph.append(pick(rng, DESIGN_PH))
    if rng.random() < 0.03:
        ph.append(pick(rng, COMPETE_PH))
    k = pick(rng, [0, 0, 1, 1, 1, 2])
    pool = NOISE + NOISE_BY_PROD.get(prod, []) * 2
    for _ in range(k):
        x = pick(rng, pool)
        if x not in ph and not (x == "修正なし" and any(FIX_RE.search(y) for y in ph)):
            ph.append(x)
    order = rng.permutation(len(ph))
    ph = [ph[i] for i in order]
    sep = pick(rng, ["、", "／", " "], [.7, .2, .1])
    return sep.join(ph), rush, hurry


# ----------------------------------------------------------------------------- base generation
SEASON = {1: 0.8, 2: 1.0, 3: 1.3, 4: 1.15, 5: 0.95, 6: 1.0, 7: 1.0, 8: 0.75, 9: 1.05, 10: 1.05,
          11: 1.1, 12: 1.15}


def gen_base(rng, custs):
    days, w = [], []
    d = START
    while d <= END:
        if is_open(d):
            days.append(d)
            x = SEASON[d.month] * (1 + 0.07 * (d - START).days / 365.0)
            x *= 1.12 if d.weekday() == 0 else 1.0
            x *= 0.08 if d.weekday() == 5 else 1.0
            w.append(x)
        d += dt.timedelta(1)
    w = np.asarray(w)
    counts = rng.multinomial(N_BASE, w / w.sum())
    quotes = []
    for d, n in zip(days, counts):
        for _ in range(int(n)):
            act = [c for c in custs if c["start"] <= d <= c["end"]]
            cw = np.asarray([c["w"] for c in act])
            cust = act[int(rng.choice(len(act), p=cw / cw.sum()))]
            ind = cust["industry"]
            repeat_p = 0.6 if cust["ctype"] in "CDG" else 0.2
            is_repeat, qty_up = False, False
            if cust["templates"] and rng.random() < repeat_p:
                t = cust["templates"][int(rng.integers(len(cust["templates"])))]
                spec = dict(t)
                is_repeat = True
                lad = QTY_LADDER[ladder_key(spec["product"])]
                u = rng.random()
                if u >= 0.6:
                    if u < 0.85 and spec["qty"] in lad:
                        j = lad.index(spec["qty"]) + pick(rng, [-1, 1])
                        j = min(max(j, 0), len(lad) - 1)
                        newq = lad[j]
                    else:
                        lo = 1 if spec["product"] in ("チラシ", "リーフレット") else 0
                        newq = pick(rng, lad[lo:-1])
                    qty_up = newq > spec["qty"]
                    spec["qty"] = newq
                if spec["product"] == "名刺" and rng.random() < 0.5:
                    spec["kinds"] = max(1, spec["kinds"] + pick(rng, [-2, -1, 1, 1, 2, 3]))
            else:
                prod = pick(rng, PRODUCTS, INDUSTRY_MIX[ind])
                spec = make_spec(rng, prod, ind)
                if (cust["ctype"] in "CDG" or rng.random() < 0.3) and len(cust["templates"]) < 5 \
                        and rng.random() < 0.55:
                    cust["templates"].append(dict(spec))
            if cust["home"] is None:
                spec["delivery"] = pick(rng, ["引取", "市内", "市外", "県外"], [.8, .08, .04, .08])
            else:
                spec["delivery"] = "引取" if rng.random() < cust["pickup_p"] else cust["home"]
            rep = rep_for(rng, cust, d)
            notes, rush, hurry = make_notes(rng, spec, is_repeat, qty_up)
            m = method_of(spec)
            if rush:
                lead = int(rng.integers(1, 3))
            elif hurry:
                lead = int(rng.integers(2, 5))
            else:
                lo, hi = {"card": (3, 7), "env": (5, 11), "od": (3, 8)}.get(m, (5, 13))
                if spec["product"] == "冊子":
                    lo, hi = 7, 16
                lead = int(rng.integers(lo, hi))
            blank_p = 0.4 if ("見積のみ" in notes or "検討中" in notes) else 0.12
            due = None if (rng.random() < blank_p and not rush) else add_bdays(d, lead)
            q = {"date": d, "cust": cust, "rep": rep, "spec": spec, "obs_spec": dict(spec), "notes": notes,
                 "due": due, "order": float(rng.random()), "kind": "base", "quirk": "clean", "qdetail": "",
                 "force_rush": False, "tax_err": None, "typo": False, "related": None}
            q["calc"] = compute(spec, notes, d, rep, cust_rate(cust, d))
            q["true_net"] = q["calc"]["net"]
            quotes.append(q)
    return quotes


# ----------------------------------------------------------------------------- quirks
QUIRK_PLAN = [("unreconcilable", 7), ("stale_price_table", 24), ("stale_customer_rate", 5),
              ("copied_customer_rate", 12), ("undocumented_discount", 18), ("unrecorded_rush", 12),
              ("rounding_inconsistency", 15), ("omitted_delivery_fee", 6), ("tax_label_error", 7),
              ("quantity_entry_error", 5), ("digit_typo", 8)]
N_REVISIONS = 18
N_DUPLICATES = 9

ALT_ROUNDERS = {
    10: [("half_up_100", lambda x, u: half_up_to(x, 100)), ("ceil_100", lambda x, u: ceil_to(x, 100)),
         ("no_rounding", lambda x, u: floor_to(x, 1))],
    100: [("floor_10", lambda x, u: floor_to(x, 10)), ("half_up_1000", lambda x, u: half_up_to(x, 1000)),
          ("ceil_100", lambda x, u: ceil_to(x, 100))],
}


def recompute(q, **kw):
    rate = kw.pop("rate", None)
    if rate is None:
        rate = cust_rate(q["cust"], q["date"])
    return compute(q["spec"], q["notes"], q["date"], q["rep"], rate, **kw)


def choose(rng, cands, k, weights=None):
    if len(cands) < k:
        raise RuntimeError("not enough candidates")
    if weights is None:
        idx = rng.choice(len(cands), size=k, replace=False)
    else:
        w = np.asarray(weights, float)
        idx = rng.choice(len(cands), size=k, replace=False, p=w / w.sum())
    return [cands[int(i)] for i in sorted(idx)]


def apply_quirks(rng, quotes):
    n = len(quotes)
    taken = [False] * n

    def free(i):
        return not taken[i]

    def mark(i, label, detail):
        taken[i] = True
        quotes[i]["quirk"] = label
        quotes[i]["qdetail"] = detail

    for label, k in QUIRK_PLAN:
        if label == "unreconcilable":
            cands = [i for i in range(n) if free(i) and quotes[i]["true_net"] >= 3000]
            for i in choose(rng, cands, k):
                q = quotes[i]
                u = rng.random()
                if u < 0.4:
                    v = floor_to(q["true_net"] * (0.35 + 0.25 * rng.random()), 1000)
                elif u < 0.8:
                    v = floor_to(q["true_net"] * (1.7 + 1.1 * rng.random()), 1000)
                else:
                    v = pick(rng, [19800, 29800, 48000, 98000, 150000])
                if rng.random() < 0.4:
                    extra = pick(rng, ["セット価格", "別紙明細参照", "年間契約分"])
                    q["notes"] = (q["notes"] + "、" + extra) if q["notes"] else extra
                q["true_net"] = int(max(v, 1000))
                mark(i, label, "amount not derivable from the spec (orig rule net %d)" % q["calc"]["net"])
        elif label == "stale_price_table":
            cands, wts = [], []
            for i in range(n):
                if not free(i):
                    continue
                q = quotes[i]
                for r in REV_DATES:
                    days = (q["date"] - r).days
                    if 0 <= days < 60:
                        alt = recompute(q, period=period_of(q["date"]) - 1)
                        if alt["net"] != q["true_net"]:
                            cands.append(i)
                            wts.append(math.exp(-days / 25.0))
            for i in choose(rng, cands, k, wts):
                q = quotes[i]
                p = period_of(q["date"]) - 1
                q["calc"] = recompute(q, period=p)
                q["true_net"] = q["calc"]["net"]
                mark(i, label, "priced with period-%d tables after revision" % p)
        elif label == "stale_customer_rate":
            cands = []
            for i in range(n):
                q = quotes[i]
                if not free(i) or len(q["cust"]["rates"]) < 2:
                    continue
                since = q["cust"]["rates"][1][0]
                if 0 <= (q["date"] - since).days < 150:
                    cands.append(i)
            for i in choose(rng, cands, k):
                q = quotes[i]
                old = q["cust"]["rates"][0][1]
                q["calc"] = recompute(q, rate=old)
                q["true_net"] = q["calc"]["net"]
                mark(i, label, "old customer rate %s used" % old)
        elif label == "copied_customer_rate":
            cands, src = [], {}
            for i in range(n):
                if not free(i):
                    continue
                q = quotes[i]
                myr = cust_rate(q["cust"], q["date"])
                for j in range(i - 1, max(-1, i - 40), -1):
                    p = quotes[j]
                    if (q["date"] - p["date"]).days > 7:
                        break
                    if p["rep"] == q["rep"] and p["cust"]["code"] != q["cust"]["code"]:
                        pr = cust_rate(p["cust"], p["date"])
                        if pr != myr:
                            alt = recompute(q, rate=pr)
                            if alt["net"] != q["true_net"]:
                                cands.append(i)
                                src[i] = (j, pr)
                        break
            for i in choose(rng, cands, k):
                q = quotes[i]
                j, pr = src[i]
                q["calc"] = recompute(q, rate=pr)
                q["true_net"] = q["calc"]["net"]
                q["copied_from_idx"] = j
                mark(i, label, "rate %s copied from previous quote for %s" % (pr, quotes[j]["cust"]["code"]))
        elif label == "undocumented_discount":
            cands = [i for i in range(n) if free(i) and quotes[i]["cust"]["ctype"] in "CDG"
                     and quotes[i]["true_net"] >= 15000]
            wts = [8.0 if "相見積" in quotes[i]["notes"] else 1.0 for i in cands]
            for i in choose(rng, cands, k, wts):
                q = quotes[i]
                net = q["true_net"]
                for _ in range(10):
                    t = pick(rng, ["floor_1000", "minus_5pct", "minus_fixed", "price_point"])
                    if t == "floor_1000":
                        v = floor_to(net, 1000)
                    elif t == "minus_5pct":
                        v = floor_to(F(net) * F(95, 100), 100)
                    elif t == "minus_fixed":
                        v = net - pick(rng, [1000, 2000, 3000])
                    else:
                        v = floor_to(net, 1000) - 200
                    if v < net:
                        break
                q["true_net"] = int(v)
                mark(i, label, "%s: %d -> %d" % (t, net, v))
        elif label == "unrecorded_rush":
            cands = [i for i in range(n) if free(i) and not RUSH_RE.search(quotes[i]["notes"])
                     and quotes[i]["due"] is not None]
            for i in choose(rng, cands, k):
                q = quotes[i]
                q["force_rush"] = True
                q["calc"] = recompute(q, force_rush=True)
                q["true_net"] = q["calc"]["net"]
                q["due"] = add_bdays(q["date"], int(rng.integers(1, 3)))
                mark(i, label, "rush surcharge applied without note")
        elif label == "rounding_inconsistency":
            cands = []
            for i in range(n):
                q = quotes[i]
                if free(i) and q["rep"] != "高橋":
                    cands.append(i)
            picked = 0
            order = rng.permutation(len(cands))
            for oi in order:
                if picked >= k:
                    break
                i = cands[int(oi)]
                q = quotes[i]
                opts = ALT_ROUNDERS[q["calc"]["unit"]]
                for o in rng.permutation(len(opts)):
                    name, fn = opts[int(o)]
                    alt = recompute(q, rounder=fn)
                    if alt["net"] != q["true_net"]:
                        q["calc"] = alt
                        q["true_net"] = alt["net"]
                        mark(i, label, "rounded with %s instead of floor-%d" % (name, q["calc"]["unit"]))
                        picked += 1
                        break
        elif label == "omitted_delivery_fee":
            cands = [i for i in range(n) if free(i) and quotes[i]["calc"]["delivery"] > 0]
            for i in choose(rng, cands, k):
                q = quotes[i]
                q["calc"] = recompute(q, skip_delivery=True)
                q["true_net"] = q["calc"]["net"]
                mark(i, label, "delivery fee %d not added" % q["calc"]["would_delivery"])
        elif label == "tax_label_error":
            cands = []
            for i in range(n):
                q = quotes[i]
                if not free(i):
                    continue
                ct = q["cust"]["ctype"]
                if q["date"] < INVOICE_DATE and ct in "CD":
                    cands.append((i, "gross_as_net"))
                elif q["date"] >= INVOICE_DATE and ct in "CDG":
                    cands.append((i, "double_tax"))
                elif ct in "KW":
                    cands.append((i, "net_as_gross"))
            for i, t in choose(rng, cands, k):
                quotes[i]["tax_err"] = t
                mark(i, label, t)
        elif label == "quantity_entry_error":
            cands = [i for i in range(n) if free(i) and quotes[i]["spec"]["product"] in
                     ("チラシ", "リーフレット", "封筒", "名刺")]
            picked = 0
            for oi in rng.permutation(len(cands)):
                if picked >= k:
                    break
                i = cands[int(oi)]
                q = quotes[i]
                obs = dict(q["spec"])
                prod = obs["product"]
                if prod == "名刺":
                    if obs["kinds"] > 1:
                        obs["kinds"] = 1
                    else:
                        obs["qty"] = obs["qty"] * 10
                else:
                    t = pick(rng, ["x10", "div10", "neighbor"])
                    lad = QTY_LADDER[ladder_key(prod)]
                    if t == "x10" and obs["qty"] * 10 <= 100000:
                        obs["qty"] *= 10
                    elif t == "div10" and obs["qty"] % 10 == 0 and obs["qty"] >= 500:
                        obs["qty"] //= 10
                    elif obs["qty"] in lad:
                        j = lad.index(obs["qty"])
                        obs["qty"] = lad[j + 1] if j + 1 < len(lad) else lad[j - 1]
                    else:
                        continue
                alt = compute(obs, q["notes"], q["date"], q["rep"], cust_rate(q["cust"], q["date"]))
                if alt["net"] == q["true_net"]:
                    continue
                q["obs_spec"] = obs
                mark(i, label, "recorded qty/kinds %d/%d, priced %d/%d" % (obs["qty"], obs["kinds"],
                                                                          q["spec"]["qty"], q["spec"]["kinds"]))
                picked += 1
        elif label == "digit_typo":
            cands = [i for i in range(n) if free(i) and quotes[i]["true_net"] >= 1000]
            for i in choose(rng, cands, k):
                quotes[i]["typo"] = True
                mark(i, label, "")  # detail filled at render time

    # re-issues (v1 superseded by v2)
    cands = [i for i in range(n) if free(i) and quotes[i]["cust"]["ctype"] in "CDG"
             and quotes[i]["spec"]["product"] != "名刺" and quotes[i]["date"] <= D(2025, 9, 10)]
    v1s = choose(rng, cands, N_REVISIONS)
    extra = []
    for t_i, i in enumerate(v1s):
        q = quotes[i]
        taken[i] = True
        q["quirk"] = "superseded_by_revision"
        v2 = {k: v for k, v in q.items()}
        v2["spec"] = dict(q["spec"])
        v2["date"] = add_bdays(q["date"], int(rng.integers(2, 9)))
        v2["kind"] = "v2"
        if v2["rep"] not in active_reps(v2["date"]):
            v2["rep"] = rep_hist_at(q["cust"], v2["date"])
        v2["order"] = float(rng.random())
        v2["related"] = q
        q["related"] = v2
        if t_i % 2 == 0:
            lad = QTY_LADDER[ladder_key(q["spec"]["product"])]
            j = lad.index(q["spec"]["qty"]) if q["spec"]["qty"] in lad else 3
            j2 = j + pick(rng, [-1, 1, 1])
            v2["spec"]["qty"] = lad[min(max(j2, 1), len(lad) - 1)]
            if v2["spec"]["qty"] == q["spec"]["qty"]:
                v2["spec"]["qty"] = lad[min(j + 1, len(lad) - 1)] if j + 1 < len(lad) else lad[j - 1]
            v2["notes"] = pick(rng, ["再見積（数量変更）", "数量変更につき再見積"]) + (
                ("、" + q["notes"]) if q["notes"] else "")
            v2["obs_spec"] = dict(v2["spec"])
            v2["calc"] = compute(v2["spec"], v2["notes"], v2["date"], v2["rep"], cust_rate(q["cust"], v2["date"]))
            v2["true_net"] = v2["calc"]["net"]
            v2["quirk"], v2["qdetail"] = "clean", "re-issue with changed quantity"
            q["qdetail"] = "re-issued with changed quantity"
        else:
            if rng.random() < 0.6:
                v2["notes"] = pick(rng, ["再見積", "再見積（お値引き）"]) + (("、" + q["notes"]) if q["notes"] else "")
            v2["obs_spec"] = dict(v2["spec"])
            v2["calc"] = compute(v2["spec"], v2["notes"], v2["date"], v2["rep"], cust_rate(q["cust"], v2["date"]))
            base_net = v2["calc"]["net"]
            unit = pick(rng, [100, 1000])
            v2["true_net"] = int(floor_to(F(base_net) * F(int(88 + rng.integers(0, 9)), 100), unit))
            v2["quirk"], v2["qdetail"] = "negotiated_revision", "price cut on re-issue: %d -> %d" % (
                base_net, v2["true_net"])
            q["qdetail"] = "re-issued with a negotiated price"
        due_lead = None if q["due"] is None else max(1, int(rng.integers(4, 12)))
        v2["due"] = None if due_lead is None else add_bdays(v2["date"], due_lead)
        extra.append(v2)
    # duplicate register entries
    cands = [i for i in range(n) if free(i)]
    for i in choose(rng, cands, N_DUPLICATES):
        q = quotes[i]
        taken[i] = True
        dup = {k: v for k, v in q.items()}
        dup["kind"] = "dup"
        dup["quirk"], dup["qdetail"] = "duplicate_entry", "same quote keyed twice"
        dup["date"] = q["date"] if rng.random() < 0.6 else add_bdays(q["date"], 1)
        dup["order"] = float(rng.random())
        dup["related"] = q
        extra.append(dup)
    return quotes + extra


# ----------------------------------------------------------------------------- rendering amounts
def typo_of(rng, amount):
    s = str(amount)
    for _ in range(20):
        t = pick(rng, ["swap", "drop0", "add0", "digit"])
        if t == "swap" and len(s) >= 3:
            j = int(rng.integers(0, len(s) - 1))
            if s[j] != s[j + 1] and not (j == 0 and s[1] == "0"):
                return int(s[:j] + s[j + 1] + s[j] + s[j + 2:]), "transposed digits"
        elif t == "drop0" and s.endswith("0"):
            return int(s[:-1]), "dropped a zero"
        elif t == "add0":
            return int(s + "0"), "extra zero"
        elif t == "digit":
            j = int(rng.integers(0, len(s) - 2)) if len(s) > 2 else 0
            dgt = int(s[j])
            nd = (dgt + pick(rng, [-1, 1])) % 10
            if j == 0 and nd == 0:
                continue
            return int(s[:j] + str(nd) + s[j + 1:]), "wrong digit"
    return int(s + "0"), "extra zero"


def render(rng, q):
    net = q["true_net"]
    ct = q["cust"]["ctype"]
    u = rng.random()
    if q["date"] < INVOICE_DATE:
        if ct in "CD":
            basis, label = "net", ("" if (q["rep"] == "鈴木" or u < 0.3) else "税抜")
        else:
            basis, label = "gross", ("" if (q["rep"] == "鈴木" or u < 0.45) else "税込")
        post = False
    else:
        basis, label = ("net", "税抜") if ct in "CDG" else ("gross", "税込")
        post = True
    gross = net + net // 10
    te = q["tax_err"]
    if te == "gross_as_net":
        amount = gross
    elif te == "double_tax":
        amount = gross
    elif te == "net_as_gross":
        amount = net
    else:
        amount = net if basis == "net" else gross
    intended = amount
    if q["typo"]:
        amount, how = typo_of(rng, amount)
        q["qdetail"] = "%s: intended %d, written %d" % (how, intended, amount)
    tax = None
    if post:
        tax = amount // 10 if label == "税抜" else (amount * 10) // 110
    q["amount"], q["label"], q["tax"], q["intended_amount"], q["basis"] = amount, label, tax, intended, basis


def written_net_of(q):
    lab, ct, amt = q["label"], q["cust"]["ctype"], q["amount"]
    if lab == "税込" or (lab == "" and ct in "GKW"):
        return int(half_up_to(F(amt) * 10 / 11, 1))
    return amt


# ----------------------------------------------------------------------------- ids
def assign_ids(rows):
    base = [r for r in rows if r["kind"] != "v2"]
    base.sort(key=lambda r: (r["date"], r["order"]))
    seq = {}
    for r in base:
        ym = r["date"].strftime("%y%m")
        seq[ym] = seq.get(ym, 0) + 1
        r["id"] = "Q%s-%03d" % (ym, seq[ym])
    for r in rows:
        if r["kind"] == "v2":
            r["id"] = r["related"]["id"] + "-R2"
    rows.sort(key=lambda r: (r["date"], r["id"]))
    return rows


# ----------------------------------------------------------------------------- public CSV
CSV_COLS = ["quote_id", "quote_date", "customer_code", "rep", "product", "size", "paper", "colors", "pages",
            "finishing", "quantity", "kinds", "delivery", "due_date", "notes", "amount", "tax_label",
            "tax_amount"]


def csv_row(r):
    s = r["obs_spec"]
    return [r["id"], r["date"].isoformat(), r["cust"]["code"], r["rep"], s["product"], s["size"], s["paper"],
            s["colors"], s["pages"] if s["pages"] else "", s["finishing"], s["qty"], s["kinds"], s["delivery"],
            r["due"].isoformat() if r["due"] else "", r["notes"], r["amount"], r["label"],
            "" if r["tax"] is None else r["tax"]]


def write_csv(path, header, rows):
    with open(path, "w", encoding="utf-8", newline="") as f:
        w = csv.writer(f, lineterminator="\n")
        w.writerow(header)
        for row in rows:
            w.writerow(row)


# ----------------------------------------------------------------------------- messy helpers
ZEN = str.maketrans("0123456789,-/.ABCDEFGHIJKLMNOPQRSTUVWXYZ()", "０１２３４５６７８９，－／．ＡＢＣＤＥＦＧＨＩＪＫＬＭＮＯＰＱＲＳＴＵＶＷＸＹＺ（）")
HANKANA = [("リーフレット", "ﾘｰﾌ"), ("チラシ", "ﾁﾗｼ"), ("マットコート", "ﾏｯﾄ"), ("コート", "ｺｰﾄ"),
           ("クラフト", "ｸﾗﾌﾄ"), ("白ケント", "白ｹﾝﾄ"), ("ヴァンヌーボ", "ｳﾞｧﾝﾇｰﾎﾞ"), ("パール紙", "ﾊﾟｰﾙ"),
           ("片面PP", "PP")]


def zen(s):
    return str(s).translate(ZEN)


def hankana(s):
    for a, b in HANKANA:
        s = s.replace(a, b)
    return s


def wareki(d, style):
    y = d.year - 2018
    if style == "dot":
        return "R%d.%d.%d" % (y, d.month, d.day)
    return "令和%d年%d月%d日" % (y, d.month, d.day)


def comma(n):
    return "{:,}".format(n)


def normalize_xlsx(path):
    with zipfile.ZipFile(path) as z:
        items = [(i.filename, z.read(i.filename)) for i in z.infolist()]
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as z:
        for name, data in items:
            if name == "docProps/core.xml":
                data = re.sub(rb"<dcterms:created([^>]*)>[^<]*</dcterms:created>",
                              rb"<dcterms:created\1>2025-10-01T09:00:00Z</dcterms:created>", data)
                data = re.sub(rb"<dcterms:modified([^>]*)>[^<]*</dcterms:modified>",
                              rb"<dcterms:modified\1>2025-10-01T09:00:00Z</dcterms:modified>", data)
            zi = zipfile.ZipInfo(name, date_time=FIXED_TS)
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            z.writestr(zi, data)
    with open(path, "wb") as f:
        f.write(buf.getvalue())


def new_wb(author):
    wb = openpyxl.Workbook()
    fixed = dt.datetime(2025, 10, 1, 9, 0, 0)
    wb.properties.created = fixed
    wb.properties.modified = fixed
    wb.properties.creator = author
    wb.properties.lastModifiedBy = author
    wb.remove(wb.active)
    return wb


THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def colors_variant(rng, c):
    v = {"4/0": ["4/0", "4C/0C", "片面カラー"], "4/4": ["4/4", "4C/4C", "両面カラー"],
         "1/0": ["1/0", "スミ1色", "1C/0"], "1/1": ["1/1", "両面スミ", "1C/1C"], "2/0": ["2/0", "2色", "2C/0"]}[c]
    return pick(rng, v, [.5, .3, .2])


def paper_variant(rng, p):
    u = rng.random()
    if u < 0.5:
        return p
    if u < 0.85:
        return p.replace("kg", "").replace("g", "")
    return hankana(p)


def product_variant(rng, p):
    v = {"チラシ": ["チラシ", "ちらし", "チラシ"], "リーフレット": ["リーフレット", "リーフ", "パンフ（折）"],
         "冊子": ["冊子", "パンフ", "冊子"], "名刺": ["名刺", "名刺", "名刺"], "封筒": ["封筒", "封筒", "封筒印刷"]}[p]
    return pick(rng, v, [.6, .25, .15])


# ----------------------------------------------------------------------------- messy 1: FY2023 ledger
def write_ledger(rng, rows, path):
    sel = [r for r in rows if D(2023, 4, 1) <= r["date"] <= D(2024, 3, 31) and r["rep"] in ("田中", "鈴木")
           and r["quirk"] != "duplicate_entry"]
    wb = new_wb("田中")
    months = [(2023, m) for m in range(4, 13)] + [(2024, m) for m in range(1, 4)]
    for (y, m) in months:
        ws = wb.create_sheet("%d月" % m)
        mr = [r for r in sel if r["date"].year == y and r["date"].month == m]
        ws["A1"] = "令和%d年度　見積台帳（営業1課）　%s月分" % (5, zen(m))
        ws.merge_cells("A1:O1")
        ws["A1"].font = Font(bold=True, size=13)
        last = (D(y, m, 28) + dt.timedelta(8)).replace(day=3)
        ws["A2"] = "作成：田中・鈴木　／　最終更新：" + wareki(last, "kanji")
        ws.merge_cells("A2:H2")
        heads1 = [("A4:B4", "見積"), ("C4:C5", "得意先"), ("D4:D5", "担当"), ("E4:I4", "仕様"),
                  ("J4:J5", "数量"), ("K4:M4", "金額"), ("N4:N5", "納期"), ("O4:O5", "備考")]
        for rng_a, txt in heads1:
            ws[rng_a.split(":")[0]] = txt
            ws.merge_cells(rng_a)
        for col, txt in zip("ABEFGHIKLM", ["No.", "日付", "品名", "サイズ", "用紙", "色", "加工", "金額", "税", "消費税"]):
            ws["%s5" % col] = txt
        for row in (4, 5):
            for col in "ABCDEFGHIJKLMNO":
                c = ws["%s%d" % (col, row)]
                c.font = Font(bold=True)
                c.alignment = Alignment(horizontal="center", vertical="center")
                c.border = BOX
        for col, wdt in zip("ABCDEFGHIJKLMNO", [11, 14, 8, 6, 12, 7, 14, 9, 10, 14, 11, 5, 8, 10, 36]):
            ws.column_dimensions[col].width = wdt
        R = 6
        month_sum = 0
        for r in mr:
            s = r["obs_spec"]
            start = R
            u = rng.random()
            qid = r["id"][1:]
            ws["A%d" % R] = zen(qid) if u < 0.1 else (r["id"] if u < 0.2 else qid)
            u = rng.random()
            if u < 0.6:
                ws["B%d" % R] = wareki(r["date"], "dot")
            elif u < 0.85:
                ws["B%d" % R] = r["date"]
                ws["B%d" % R].number_format = '[$-ja-JP-x-gannen]ggge"年"m"月"d"日";@'
            else:
                ws["B%d" % R] = wareki(r["date"], "kanji")
            ws["C%d" % R] = zen(r["cust"]["code"]) if rng.random() < 0.05 else r["cust"]["code"]
            ws["D%d" % R] = r["rep"]
            ws["E%d" % R] = product_variant(rng, s["product"])
            ws["F%d" % R] = zen(s["size"]) if rng.random() < 0.15 else s["size"]
            ws["G%d" % R] = paper_variant(rng, s["paper"])
            ws["H%d" % R] = colors_variant(rng, s["colors"])
            fin = s["finishing"]
            if s["product"] == "冊子":
                fin = "中綴じ%dP" % s["pages"] if rng.random() < 0.7 else "%dページ 中綴" % s["pages"]
            ws["I%d" % R] = fin
            unit = {"冊子": "部", "名刺": "枚", "封筒": "枚"}.get(s["product"], "枚")
            u = rng.random()
            if s["product"] == "名刺":
                ws["J%d" % R] = ("%d枚×%d名" % (s["qty"], s["kinds"])) if u < 0.7 else ("%d名×各%d" % (s["kinds"], s["qty"]))
            elif s["kinds"] > 1:
                ws["J%d" % R] = "%s×%d種" % (comma(s["qty"]), s["kinds"])
            elif u < 0.5:
                ws["J%d" % R] = s["qty"]
            elif u < 0.65:
                ws["J%d" % R] = zen(comma(s["qty"])) + unit
            else:
                ws["J%d" % R] = comma(s["qty"]) + unit
            amt = r["amount"]
            subrows = []
            if r["basis"] == "net" and r["label"] != "税込" and r["tax_err"] is None and not r["typo"]:
                if r["quirk"] == "undocumented_discount":
                    pre = int(re.search(r"(\d+) ->", r["qdetail"]).group(1))
                    subrows.append(("　└出精値引", r["true_net"] - pre))
                    amt = pre
                else:
                    c = r["calc"]
                    if c["delivery"] > 0 and rng.random() < 0.4 and c["rep_adj"] == 0 \
                            and r["quirk"] != "unreconcilable":
                        subrows.append(("　└送料", c["delivery"]))
                        amt = r["amount"] - c["delivery"]
                    if r["quirk"] == "clean" and rng.random() < 0.05:
                        raw = math.floor(c["pre_round"])
                        if raw != c["sub"]:
                            subrows.append(("　└端数調整", c["sub"] - raw))
                            amt = amt - (c["sub"] - raw)
            u = rng.random()
            if subrows or u < 0.7:
                ws["K%d" % R] = amt
                ws["K%d" % R].number_format = "#,##0"
                month_sum += amt
            elif u < 0.85:
                ws["K%d" % R] = "¥" + comma(amt)
            elif u < 0.95:
                ws["K%d" % R] = comma(amt) + "円"
            else:
                ws["K%d" % R] = zen(comma(amt))
            lab = r["label"]
            ws["L%d" % R] = {"税抜": pick(rng, ["外", "税抜", "抜"]), "税込": pick(rng, ["込", "税込", "内"]),
                             "": ""}[lab]
            if r["tax"] is not None:
                ws["M%d" % R] = r["tax"]
            if r["due"]:
                u = rng.random()
                ws["N%d" % R] = ("%d/%d" % (r["due"].month, r["due"].day)) if u < 0.6 else (
                    "%d月%d日" % (r["due"].month, r["due"].day))
            else:
                ws["N%d" % R] = pick(rng, ["", "未定", "相談"])
            ws["O%d" % R] = r["notes"]
            for lab_s, v in subrows:
                R += 1
                ws["E%d" % R] = lab_s
                ws["K%d" % R] = v
                ws["K%d" % R].number_format = "#,##0;[Red]-#,##0"
                month_sum += v
            r["ref_ledger"] = "%s!R%d" % (ws.title, start) + ("-R%d" % R if R > start else "")
            R += 1
        R += 1
        ws["E%d" % R] = "月計"
        ws["K%d" % R] = month_sum
        ws["K%d" % R].number_format = "#,##0"
        ws["E%d" % R].font = Font(bold=True)
        ws["K%d" % R].font = Font(bold=True)
        ws["L%d" % R] = "※税込・税抜混在"
        ws.freeze_panes = "A6"
    wb.save(path)
    normalize_xlsx(path)


# ----------------------------------------------------------------------------- messy 2: filed copies
def spec_text(s):
    if s["product"] == "冊子":
        t = "%s　%dP　中綴じ　%s　%s" % (s["size"], s["pages"], s["paper"], s["colors"])
    elif s["product"] == "名刺":
        t = "91×55mm　%s　%s" % (s["paper"], s["colors"]) + ("　" + s["finishing"] if s["finishing"] else "") + \
            "　各%d枚" % s["qty"]
    else:
        t = "%s　%s　%s" % (s["size"], s["paper"], s["colors"]) + ("　" + s["finishing"] if s["finishing"] else "")
        if s["kinds"] > 1:
            t += "　%d種" % s["kinds"]
    return t


def write_copies(rng, rows, path):
    sel = [r for r in rows if D(2024, 10, 1) <= r["date"] <= D(2025, 3, 31) and r["cust"]["ctype"] in "CG"
           and r["quirk"] != "duplicate_entry"]
    sel = [r for r in sel if rng.random() < 0.45]
    wb = new_wb("事務")
    ws = wb.create_sheet("見積書控え")
    for col, wdt in zip("ABCDEFGH", [2, 14, 40, 10, 6, 10, 14, 14]):
        ws.column_dimensions[col].width = wdt
    R = 1
    for r in sel:
        s = r["obs_spec"]
        start = R
        ws["B%d" % R] = "御　見　積　書"
        ws.merge_cells("B%d:H%d" % (R, R))
        ws["B%d" % R].font = Font(bold=True, size=16)
        ws["B%d" % R].alignment = Alignment(horizontal="center")
        ws["G%d" % (R + 1)] = "No."
        ws["H%d" % (R + 1)] = r["id"][1:]
        ws["G%d" % (R + 2)] = "見積日"
        ws["H%d" % (R + 2)] = wareki(r["date"], "kanji")
        ws["B%d" % (R + 3)] = "%s　御中" % r["cust"]["code"]
        ws.merge_cells("B%d:D%d" % (R + 3, R + 3))
        ws["B%d" % (R + 4)] = "下記のとおり御見積申し上げます。"
        amt = r["intended_amount"]
        tax = amt // 10
        ws["B%d" % (R + 5)] = "御見積金額"
        ws["C%d" % (R + 5)] = "¥%s-" % comma(amt + tax)
        ws["D%d" % (R + 5)] = "（消費税込）"
        ws["C%d" % (R + 5)].font = Font(bold=True, size=13, underline="single")
        hr = R + 7
        for col, txt in zip("BCDEFG", ["品名", "仕様", "数量", "単位", "単価", "金額"]):
            ws["%s%d" % (col, hr)] = txt
            ws["%s%d" % (col, hr)].border = BOX
            ws["%s%d" % (col, hr)].font = Font(bold=True)
        c = r["calc"]
        dlv = c["delivery"] if r["quirk"] != "unreconcilable" else 0
        main = amt - dlv
        if r["quirk"] == "tax_label_error":
            main, dlv = amt, 0
        L = hr + 1
        ws["B%d" % L] = s["product"]
        ws["C%d" % L] = spec_text(s)
        if s["product"] == "名刺":
            qty, unit = s["kinds"], "名"
        else:
            qty, unit = s["qty"] * s["kinds"], {"冊子": "部"}.get(s["product"], "枚")
        up = F(main, qty)
        if rng.random() < 0.6 and (up * 100).denominator == 1:
            ws["D%d" % L], ws["E%d" % L] = qty, unit
            ws["F%d" % L] = float(up)
            ws["F%d" % L].number_format = "#,##0.00"
        else:
            ws["D%d" % L], ws["E%d" % L], ws["F%d" % L] = 1, "式", ""
            ws["C%d" % L] = ws["C%d" % L].value + "　%s%s" % (comma(qty), unit)
        ws["G%d" % L] = main
        ws["G%d" % L].number_format = "#,##0"
        if dlv > 0:
            L += 1
            ws["B%d" % L], ws["C%d" % L], ws["D%d" % L], ws["E%d" % L] = "配送料", s["delivery"], 1, "式"
            ws["F%d" % L], ws["G%d" % L] = dlv, dlv
        for k in range(hr + 1, hr + 3):
            for col in "BCDEFG":
                ws["%s%d" % (col, k)].border = BOX
        T0 = hr + 3
        for k, (lab, v) in enumerate([("小計", amt), ("消費税（10%）", tax), ("合計", amt + tax)]):
            ws["F%d" % (T0 + k)] = lab
            ws["G%d" % (T0 + k)] = v
            ws["G%d" % (T0 + k)].number_format = "#,##0"
        ws["B%d" % (T0 + 3)] = "納期：" + (wareki(r["due"], "kanji") if r["due"] else "別途ご相談")
        ws["B%d" % (T0 + 4)] = "納品：" + {"引取": "ご来社引取", "市内": "市内配達", "市外": "市外配達",
                                          "県外": "宅配便発送"}[s["delivery"]]
        ws["B%d" % (T0 + 5)] = "備考：" + (r["notes"] or "")
        ws["B%d" % (T0 + 6)] = "有効期限：見積日より1ヶ月"
        ws["G%d" % (T0 + 6)] = "担当：" + r["rep"]
        end = T0 + 6
        r["ref_copies"] = "%s!R%d-R%d" % (ws.title, start, end)
        ws.row_breaks.append(Break(id=end + 1))
        R = end + 3
    ws.print_area = "A1:H%d" % R
    wb.save(path)
    normalize_xlsx(path)


# ----------------------------------------------------------------------------- messy 3: rep memo
def memo_qty(rng, s):
    if s["product"] == "名刺":
        return pick(rng, ["%d×%d名" % (s["qty"], s["kinds"]), "%d名 各%d" % (s["kinds"], s["qty"])])
    q = s["qty"]
    if q >= 10000 and q % 5000 == 0 and rng.random() < 0.7:
        txt = ("%g万" % (q / 10000))
    elif q >= 1000 and q % 1000 == 0 and rng.random() < 0.6:
        txt = "%d千" % (q // 1000)
    elif rng.random() < 0.3:
        txt = zen(str(q))
    else:
        txt = comma(q)
    if s["kinds"] > 1:
        txt += "×%d種" % s["kinds"]
    if s["product"] == "冊子":
        txt += "部"
    return txt


def memo_spec(s):
    p = hankana(s["product"])
    col = {"4/0": "片ｶﾗｰ", "4/4": "両ｶﾗｰ", "1/0": "片ｽﾐ", "1/1": "両ｽﾐ", "2/0": "2色"}[s["colors"]]
    paper = hankana(s["paper"]).replace("kg", "").replace("g", "")
    fin = {"三つ折": "3折", "二つ折": "2折", "Z折": "Z折", "片面PP": "PP", "中綴じ": "中綴", "角丸": "角丸",
           "片面PP+角丸": "PP角丸", "": ""}[s["finishing"]]
    if s["product"] == "冊子":
        return " ".join(x for x in [p, s["size"], "%dP" % s["pages"], paper, col, fin] if x)
    if s["product"] == "名刺":
        return " ".join(x for x in [p, paper, col, fin] if x)
    return " ".join(x for x in [p, s["size"], paper, col, fin] if x)


def write_memo(rng, rows, path):
    sel = [r for r in rows if r["rep"] == "高橋" and D(2024, 4, 1) <= r["date"] <= D(2025, 9, 30)
           and r["quirk"] != "duplicate_entry"]
    wb = new_wb("高橋")
    ws = wb.create_sheet("見積メモ")
    ws["A1"] = "見積メモ（高橋）"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "※金額は見積書記載どおり"
    for col, txt in zip("ABCDEFGH", ["日付", "客先", "内容", "数量", "金額", "税", "結果", "メモ"]):
        ws["%s4" % col] = txt
        ws["%s4" % col].font = Font(bold=True)
    for col, wdt in zip("ABCDEFGH", [12, 8, 30, 12, 12, 5, 6, 30]):
        ws.column_dimensions[col].width = wdt
    R = 5
    cur = None
    for r in sel:
        ym = (r["date"].year, r["date"].month)
        if ym != cur:
            if cur is not None:
                R += 1
            ws["A%d" % R] = "【%d年%d月】" % ym
            ws["A%d" % R].font = Font(bold=True)
            R += 1
            cur = ym
        s = r["obs_spec"]
        d = r["date"]
        u = rng.random()
        if u < 0.6:
            ws["A%d" % R] = d
            ws["A%d" % R].number_format = "yyyy/m/d"
        elif u < 0.8:
            ws["A%d" % R] = "%d/%d/%d" % (d.year, d.month, d.day)
        elif u < 0.95:
            ws["A%d" % R] = "%d/%d" % (d.month, d.day)
        else:
            ws["A%d" % R] = wareki(d, "dot")
        code = r["cust"]["code"]
        u = rng.random()
        if code[0] in "CDGK" and u < 0.35:
            code = code[0] + str(int(code[1:]))
        elif code[0] in "CDGK" and u < 0.45:
            code = code[0].lower() + "-" + code[1:]
        elif u < 0.5:
            code = zen(code)
        ws["B%d" % R] = code
        ws["C%d" % R] = memo_spec(s) + ("（再）" if r["kind"] == "v2" else "")
        ws["D%d" % R] = memo_qty(rng, s)
        amt = r["intended_amount"]
        u = rng.random()
        if u < 0.6:
            ws["E%d" % R] = amt
            ws["E%d" % R].number_format = "#,##0"
        elif u < 0.75:
            ws["E%d" % R] = "¥" + comma(amt)
        elif u < 0.9 and amt % 10 == 0:
            ws["E%d" % R] = "%s千" % ("%g" % (amt / 1000))
        else:
            ws["E%d" % R] = zen(comma(amt))
        ws["F%d" % R] = {"税抜": pick(rng, ["別", "外税", "抜"]), "税込": pick(rng, ["込", "内"]), "": ""}[r["label"]]
        res = pick(rng, ["受注", "失注", "保留", "", "取消"], [.55, .25, .1, .07, .03])
        if r["quirk"] == "superseded_by_revision":
            res = "→再見積"
        ws["G%d" % R] = res
        memo = ""
        if r["quirk"] == "undocumented_discount" and rng.random() < 0.6:
            memo = pick(rng, ["社長OK", "値引済", "社長了承済み"])
        elif r["notes"]:
            memo = re.split("[、／ ]", r["notes"])[0]
        if rng.random() < 0.08:
            memo = (memo + " " if memo else "") + pick(rng, ["要フォロー", "TEL済", "再来週返事", "前回より安く"])
        ws["H%d" % R] = memo
        if res == "取消":
            for col in "ABCDEFGH":
                ws["%s%d" % (col, R)].font = Font(strike=True, color="FF808080")
        r["ref_memo"] = "%s!R%d" % (ws.title, R)
        R += 1
    ws2 = wb.create_sheet("集計")
    ws2["A1"] = "月別件数（手入力）"
    months = sorted({(r["date"].year, r["date"].month) for r in sel})
    for k, ym in enumerate(months, 2):
        ws2["A%d" % k] = "%d/%d" % ym
        ws2["B%d" % k] = sum(1 for r in sel if (r["date"].year, r["date"].month) == ym)
    wb.save(path)
    normalize_xlsx(path)


# ----------------------------------------------------------------------------- truth
def fmt(x):
    if x is None:
        return ""
    if isinstance(x, F):
        if x.denominator == 1:
            return str(x.numerator)
        return ("%.2f" % float(x)).rstrip("0").rstrip(".")
    return str(x)


TRUTH_COLS = (["quote_id", "quote_date", "customer_code", "rep", "product", "method", "period", "rate_applied"]
              + ["c_" + k for k in COMP_KEYS]
              + ["base_total", "after_rate", "min_applied", "after_min", "rush"]
              + ["x_" + k for k in EXTRA_KEYS]
              + ["pre_round", "round_unit", "sub_rounded", "delivery_fee", "rep_adjust", "true_net", "rule_net",
                 "written_net", "written_amount", "written_tax_label", "written_tax_amount", "quirk", "quirk_detail",
                 "related_quote_id", "priced_qty", "priced_kinds", "messy_ledger_ref", "messy_copies_ref",
                 "messy_memo_ref", "messy_amount_note"])


def truth_row(r):
    c = r["calc"]
    unrec = r["quirk"] == "unreconcilable"
    out = [r["id"], r["date"].isoformat(), r["cust"]["code"], r["rep"], r["spec"]["product"], c["method"],
           c["period"], fmt(c["rate"])]
    if unrec:
        out += [""] * (len(COMP_KEYS) + 5 + len(EXTRA_KEYS) + 5)
    else:
        out += [fmt(c["comp"][k]) for k in COMP_KEYS]
        out += [fmt(c["base"]), fmt(c["after_rate"]), int(c["min_applied"]), fmt(c["after_min"]), fmt(c["rush"])]
        out += [fmt(c["extras"][k]) for k in EXTRA_KEYS]
        out += [fmt(c["pre_round"]), c["unit"], c["sub"], c["delivery"], c["rep_adj"]]
    note = ""
    if r["typo"] and (r.get("ref_copies") or r.get("ref_memo")):
        note = "copies/memo show intended amount %d" % r["intended_amount"]
    out += [r["true_net"], r["rule_net"], r["written_net"], r["amount"], r["label"],
            "" if r["tax"] is None else r["tax"], r["quirk"], r["qdetail"],
            r["related"]["id"] if r["related"] is not None else "",
            r["spec"]["qty"], r["spec"]["kinds"],
            r.get("ref_ledger", ""), r.get("ref_copies", ""), r.get("ref_memo", ""), note]
    return out


# ----------------------------------------------------------------------------- docs
SCHEMA_MD = """# quotes.csv — 列定義 / Column definitions

見積台帳から書き出した見積の一覧です。1行が1通の見積書に対応します。
One row per quote document as recorded in the shop's quote register.

- 文字コード / Encoding: UTF-8, LF, comma-separated, header row
- 並び順 / Order: `quote_date`, then `quote_id`
- 行数 / Rows: {nrows}

| column | 日本語 | English |
|---|---|---|
| `quote_id` | 見積番号。再発行した見積は元番号の末尾に `-R2` が付く | Quote number; a re-issued quote keeps the original number with suffix `-R2` |
| `quote_date` | 見積日 (YYYY-MM-DD) | Date the quote was issued |
| `customer_code` | 得意先コード。C=法人、D=広告代理店・制作会社、G=学校・団体、K=個人、`一般`=口座のない都度客 | Customer account code. C = company, D = agency / design studio, G = school / association, K = private individual, `一般` = walk-in without an account |
| `rep` | 担当者 | Sales rep who issued the quote |
| `product` | 品目（チラシ／リーフレット／冊子／名刺／封筒） | Product category (flyer / folded leaflet / saddle-stitched booklet / business card / envelope) |
| `size` | 仕上りサイズ（名刺は 91x55 mm、封筒は封筒規格） | Finished size (cards: 91x55 mm; envelopes: envelope standard) |
| `paper` | 用紙（銘柄と連量／坪量） | Paper stock as written |
| `colors` | 色数 表/裏（例: 4/0 = 表4色・裏なし） | Ink colours front/back (e.g. 4/0 = four colours front, none on back) |
| `pages` | ページ数（冊子のみ。その他は空欄） | Page count (booklets only; blank otherwise) |
| `finishing` | 加工（記載どおり） | Finishing as written |
| `quantity` | 数量。1種類あたり（名刺は1名あたりの枚数、冊子は部数） | Quantity per kind (cards: cards per person; booklets: copies) |
| `kinds` | 種類数（名刺は人数、その他は刷り分けの版数） | Number of kinds (cards: number of people; others: number of versions) |
| `delivery` | 納品方法・地域（引取／市内／市外／県外） | Delivery method / area (customer pick-up, in-city, out-of-city, out-of-prefecture) |
| `due_date` | 希望納期（空欄は未定） | Requested delivery date (blank = not set) |
| `notes` | 備考（自由記述、記載どおり） | Free-text remarks as written |
| `amount` | 見積金額（円、記載どおり） | Quoted amount in JPY, as written |
| `tax_label` | 税表示（記載どおり。空欄は記載なし） | Tax label as written; blank = no label written |
| `tax_amount` | 消費税額（記載がある場合のみ） | Consumption-tax amount, only where written |

## messy/

同じ見積の一部を、社内の別の記録から取り出したものです。書式・表記はそれぞれの記録のままです。
`quotes.csv` との対応表は付属していません。

Alternative records of subsets of the same quotes, kept in their original layouts and notations.
No mapping to `quotes.csv` is provided.

| file | 内容 | Contents |
|---|---|---|
| `messy/見積台帳_R5年度_営業1課.xlsx` | 年度の見積台帳（月別シート） | Fiscal-year quote ledger, one sheet per month |
| `messy/見積書控え_2024下期.xlsx` | 発行した見積書の控え（印刷レイアウト） | File copies of issued quote documents (print layout) |
| `messy/見積メモ_高橋.xlsx` | 担当者個人の見積メモ | One rep's personal quote notes |
"""


def rules_md(rows):
    L = []
    a = L.append
    a("# SEALED — Hidden pricing rules and quirks: print shop dataset\n")
    a("Generator: `generator.py` (seed %d). Rows in quotes.csv: %d. This file is generated by the generator from the "
      "same constants it prices with.\n" % (SEED, len(rows)))
    a("## 0. Order of operations (per quote)\n")
    a("1. Pick production **method**: 名刺 → table pricing (`card`); 封筒 → small offset (`env`); 冊子 → on-demand (`od`) "
      "if copies < 300 else `offset`; チラシ/リーフレット → `od` if quantity-per-kind < 1,000 else `offset`. "
      "The product name チラシ vs リーフレット does not matter; the `finishing` field does.")
    a("2. Sum production components (tables of the **price period** of the quote date) → `base`.")
    a("3. `after_rate = base × customer rate` (rate in force on the quote date).")
    a("4. **Minimum charge**: `after_min = max(after_rate, MIN)` — *not applied to 名刺*.")
    a("5. **特急** surcharge: +25% of `after_min`, only if notes contain the literal `特急` (`至急`, `急ぎ`, `お急ぎ` do NOT trigger).")
    a("6. Note-triggered flat **extras** (not multiplied by the customer rate): see §4.")
    a("7. `pre_round = after_min + rush + extras`; `sub = floor(pre_round / U) × U`, U = 10 yen before 2024-07-01, "
      "100 yen from 2024-07-01.")
    a("8. **Delivery** fee by area band, decided on `sub` (§5), added after rounding.")
    a("9. **Rep habit**: if rep is 高橋 and `sub + delivery ≥ 10,000`, the total is floored to a multiple of 500.")
    a("10. `net` = pre-tax net. Tax presentation (§6) turns it into the written amount.\n")
    a("## 1. Price periods (revisions)\n")
    a("| period | from | what changed |\n|---|---|---|")
    a("| P0 | 2022-10-01 | initial tables |")
    a("| P1 | 2023-04-01 | paper ×1.15 (rounded to 0.1 yen/sheet), envelope stock ×1.12 (rounded to 10 yen/box) |")
    a("| P2 | 2024-07-01 | paper ×1.08 again; envelopes ×1.06; plate 2,500→2,800; impression base 2,000→2,200 and step 700→750; "
      "cutting 800→1,000; OD click 30→32 (colour) and 8→9 (mono); minimum 5,000→5,500; **rounding unit 10→100**; "
      "data-fix fee 3,000→3,500; binding setup 3,000→3,500 and per-copy 10/15→12/18; PP 15→16 per A4-eq; "
      "envelope plate 1,800→2,000 and env impression 2,000/600→2,200/650; parcel box 1,300→1,500; mail 520→600 |")
    a("| P3 | 2025-04-01 | 名刺 only: first box +200/name, additional box +100; envelope stock ×1.05 |\n")
    a("## 2. Production components\n")
    a("### 2a. Offset flat (チラシ/リーフレット, qty ≥ 1,000)")
    a("- Press sheet = 半裁 (half of a full sheet). Up-count: A3 2, A4 4, A5 8 (菊判 family); B4 4, B5 8 (四六判 family).")
    a("- 通し `ts = ceil(qty / up)` per kind. Plates per kind = front colours + back colours.")
    a("- 版代 = plate price × plates × kinds.")
    a("- 印刷代 = plates × kinds × (imp_base + imp_add × (ceil(ts/1000) − 1)).")
    a("- 用紙: press sheets = kinds × (ts + 30 × plates_per_kind + ceil(2% × ts)); full sheets = ceil(press/2); "
      "**billed in packs of 250 full sheets** (ceil to 250); × price per full sheet (菊 for A sizes, 四六 for B sizes).")
    a("- 断裁 flat per job.\n")
    a("### 2b. On-demand flat (qty < 1,000)")
    a("- SRA3 up-count: A3 1, A4 2, A5 4, B4 1, B5 2. Sheets S = kinds × ceil(qty/up); paper = (S + max(3, ceil(3%·S))) × OD sheet price.")
    a("- Clicks on sides = S × (2 if back colours > 0 else 1). **Marginal tiers**: sides 1–300 at 100%, 301–1,000 at 80%, >1,000 at 60%. "
      "Colour click if front colours = 4, else mono.")
    a("- データ処理料 1,500 per kind, **waived if notes contain `前回同様` or `リピート`**.")
    a("- 断裁 flat per job.\n")
    a("### 2c. Finishing (flat products, either method)")
    a("- 二つ折: 2,000 + 1,200 × ceil(pieces/1000). 三つ折 / Z折: 2,500 + 1,600 × ceil(pieces/1000). (pieces = qty × kinds; no revision.)")
    a("- 片面PP: 3,000 + per-A4-equivalent price × A4eq × pieces (A3 2, A4 1, A5 0.5, B4 1.5, B5 0.75).\n")
    a("### 2d. 冊子 (中綴じ)")
    a("- OD (< 300 copies): SRA3 sheets per copy = pages/4 (A4) or pages/8 (A5); sides = 2 × sheets; same click tiers/paper/data fee as 2b (data fee once per job).")
    a("- Offset: A4 → 8 pages per full 台, A5 → 16. full = pages // per_台; a remaining half-台 is work-and-turn. "
      "Plates = full × (front+back) + half × front. Impression per plate = imp_base + imp_add × (ceil(copies/1000) − 1). "
      "Press sheets = full × copies + half × ceil(copies/2) + 30 × plates + ceil(2%·net); full sheets = ceil(press/2), packs of 250; 菊 price.")
    a("- 製本 = setup + per-copy (≤16 pages small rate, >16 pages large rate). No separate cutting.\n")
    a("### 2e. 封筒")
    a("- Envelope stock sold by **box** (長3/洋長3 1,000 per box; 角2 500 per box): ceil(qty × kinds / box) × box price. No spoilage charged.")
    a("- Plates = env plate × colours × kinds. Impressions per colour per kind = env_base + env_add × (ceil(qty/1000) − 1); **角2 × 1.5**.\n")
    a("### 2f. 名刺")
    a("- Per person: first box (100 cards) price by colours + additional-box price × (boxes − 1); special papers "
      "(ヴァンヌーボ195kg, パール紙180kg) +300 per box; 片面PP +600 per box; 角丸 +500 per person.")
    a("- **Step discount by person index**: persons 1–5 at 100%, 6–10 at 90%, 11+ at 80%.")
    a("- No minimum charge.\n")
    a("## 3. Customer rates\n")
    a("| code | rate history | industry |\n|---|---|---|")
    for c in CUSTS:
        hist = "; ".join("%s from %s" % (v, s.isoformat()) for s, v in c["rates"])
        a("| %s | %s | %s |" % (c["code"], hist, c["industry"]))
    a("")
    a("## 4. Note-triggered extras (regex on `notes`; added after the minimum, not rate-multiplied)\n")
    a("| trigger | applies to | fee |\n|---|---|---|")
    a("| `特急` | all | +25% of after_min (rush) |")
    a("| `データ修正` / `文字修正` / `修正あり` (NOT `修正なし`, NOT `肩書変更あり`) | all | 3,000 (P2+: 3,500); 名刺 1,000 |")
    a("| `特色` / `DIC` | offset and envelopes only (OD notes like `DIC近似色で` get nothing) | 4,000 |")
    a("| `色校` | offset and OD only | offset 8,000; OD 1,500 |")
    a("| `帯掛` | all but 名刺 | 400 × ceil(pieces/1000) |")
    a("| `新規デザイン` / `新規作成` | 名刺 | 3,000 + 500 × (persons − 1) |")
    a("| `前回同様` / `リピート` | OD | waives the data-processing fee (§2b) |\n")
    a("## 5. Delivery bands\n")
    a("- 引取: 0. 市内: 800, **free if sub ≥ 10,000**. 市外: 1,500, **free if sub ≥ 30,000**.")
    a("- 県外: weight-based — weight = pieces × area × g/m² (envelopes: fixed grams each; booklets: copies × pages/2 leaves). "
      "≤ 1 kg → mail rate (520; P2+ 600); else parcel boxes = ceil(kg / 15) × box rate (1,300; P2+ 1,500).\n")
    a("## 6. Tax presentation\n")
    a("- Before 2023-10-01 (pre-invoice template, no tax column): C/D customers → amount is **net**, label `税抜` or blank; "
      "G/K/一般 → amount is **gross** = net + floor(net×10%), label `税込` or blank. "
      "**A blank label therefore means net for C/D but gross for G/K/一般.** Rep 鈴木 never wrote a label; others left it blank ~30–45% of the time.")
    a("- From 2023-10-01 (invoice template): C/D/G → net, label `税抜`, tax_amount = floor(net×10%); "
      "K/一般 → gross, label `税込`, tax_amount = floor(gross×10/110) (内税).\n")
    a("## 7. Reps\n")
    a("田中 (whole period), 鈴木 (until 2025-03-31; customers then move to 渡辺 or 高橋), 高橋 (whole period; 500-yen floor habit), "
      "渡辺 (from 2023-06-01; took over part of 田中's accounts). Walk-ins are served by counter staff. ~4% of quotes are covered by another rep.\n")
    a("## 8. Price tables\n")
    for p, T in enumerate(TABLES):
        a("### P%d" % p)
        a("- paper per full sheet 菊: " + ", ".join("%s %s" % (k, fmt(v)) for k, v in T["paper_kiku"].items()))
        a("- paper per full sheet 四六 (=菊×1.38): " + ", ".join("%s %s" % (k, fmt(v)) for k, v in T["paper_46"].items()))
        a("- OD paper per SRA3 sheet (=菊×0.45): " + ", ".join("%s %s" % (k, fmt(v)) for k, v in T["paper_od"].items()))
        a("- envelope box: " + ", ".join("%s/%s %d" % (k[0], k[1], v) for k, v in T["env_box"].items()))
        a("- 名刺 first box: %s; additional box: %s" % (T["card_first"], T["card_add"]))
        scal = {k: v for k, v in T.items() if isinstance(v, int)}
        a("- scalars: " + ", ".join("%s=%d" % (k, v) for k, v in scal.items()))
        a("")
    a("## 9. Human quirks (labelled in truth.csv `quirk`)\n")
    a("`rule_net` = the rules above applied to the observable fields exactly as recorded in quotes.csv. "
      "`true_net` = the pre-tax net the shop actually meant to quote (after the quirk, before clerical typos). "
      "`written_net` = net implied by the written amount and label (blank label read per §6).\n")
    desc = {
        "clean": "Priced by the rules.",
        "stale_price_table": "Within 60 days after a revision, the rep used the previous period's whole table (incl. its rounding/minimum).",
        "stale_customer_rate": "Customer's rate changed (C003 2024-01-01, D02 2024-07-01) but the old rate was used.",
        "copied_customer_rate": "Quote copied from the same rep's previous sheet for another customer; that customer's rate stayed.",
        "undocumented_discount": "Manager discount with no note: floor to 1,000, −5% (floor 100), −1,000/2,000/3,000, or x,800 price point.",
        "unrecorded_rush": "Rush +25% applied but `特急` not written; due date set 1–2 business days out.",
        "rounding_inconsistency": "Different rounding than the period rule (half-up/ceil to 100, no rounding, floor 10, half-up 1,000).",
        "omitted_delivery_fee": "Delivery fee forgotten.",
        "tax_label_error": "gross_as_net: pre-invoice C/D quote written gross with net label; double_tax: post-invoice net-basis quote "
                           "written gross and taxed again; net_as_gross: K/一般 quote written net under a gross reading.",
        "quantity_entry_error": "Register quantity/kinds differs from what was priced (`priced_qty`/`priced_kinds`).",
        "digit_typo": "Amount mistyped in the register (transposed/dropped/extra/wrong digit). The filed copy and rep memo keep the intended amount.",
        "duplicate_entry": "Same quote keyed twice under a new number (same or next business day). Not present in messy files.",
        "superseded_by_revision": "v1 of a re-issued quote (priced by the rules); see related_quote_id.",
        "negotiated_revision": "-R2 re-issue at a negotiated lower price (88–96% of rule, floored to 100 or 1,000).",
        "unreconcilable": "Amount bears no derivable relation to the spec (package deals etc.); components blank.",
    }
    a("| quirk | count | definition |\n|---|---|---|")
    counts = {}
    for r in rows:
        counts[r["quirk"]] = counts.get(r["quirk"], 0) + 1
    for k in desc:
        a("| %s | %d | %s |" % (k, counts.get(k, 0), desc[k]))
    nq = sum(v for k, v in counts.items() if k != "clean")
    a("\nNon-clean rows: %d of %d (%.1f%%). Re-issue v2 rows with only a quantity change are `clean`.\n" % (
        nq, len(rows), 100.0 * nq / len(rows)))
    a("## 10. Messy workbooks (public `messy/`)\n")
    a("- `見積台帳_R5年度_営業1課.xlsx`: FY2023 (2023-04-01..2024-03-31), reps 田中/鈴木; per-month sheets, merged 2-row header, "
      "和暦 strings or dates with 和暦 format, 全角 digits, amount strings, monthly total row (mixes 税込/税抜). Sub-rows: "
      "`└出精値引` (reveals undocumented discounts), `└送料`, `└端数調整`; main + sub-rows = written amount. Amounts equal the register (typos included).")
    a("- `見積書控え_2024下期.xlsx`: 2024-10-01..2025-03-31, C/G customers, ~45% sample; printed-quote blocks stacked with page breaks; "
      "lines split product vs 配送料; shows the **intended** amount (digit typos corrected).")
    a("- `見積メモ_高橋.xlsx`: rep 高橋, 2024-04-01..2025-09-30; no quote ids; 半角カナ spec text, mixed date forms (some without year), "
      "千/万 quantities, amounts in several notations, outcome column, strikethrough cancelled rows; intended amounts; "
      "memo hints like `社長OK` on some discounts.")
    a("- truth.csv `messy_*_ref` = `sheet!Rstart[-Rend]` of each quote's block/row.\n")
    return "\n".join(L) + "\n"


# ----------------------------------------------------------------------------- main
CUSTS = []


def sha256_file(p):
    h = hashlib.sha256()
    with open(p, "rb") as f:
        h.update(f.read())
    return h.hexdigest()


def main():
    global CUSTS
    rng = np.random.default_rng(SEED)
    CUSTS = build_customers(rng)
    quotes = gen_base(rng, CUSTS)
    rows = apply_quirks(np.random.default_rng(SEED + 1), quotes)
    rows = assign_ids(rows)
    rrng = np.random.default_rng(SEED + 2)
    # render originals first so duplicates copy them exactly
    for r in rows:
        if r["kind"] != "dup":
            render(rrng, r)
    for r in rows:
        if r["kind"] == "dup":
            o = r["related"]
            for k in ("amount", "label", "tax", "intended_amount", "basis"):
                r[k] = o[k]
            r["obs_spec"], r["notes"], r["due"], r["rep"] = o["obs_spec"], o["notes"], o["due"], o["rep"]
    for r in rows:
        r["rule_net"] = compute(r["obs_spec"], r["notes"], r["date"], r["rep"], cust_rate(r["cust"], r["date"]))["net"]
        r["written_net"] = written_net_of(r)

    os.makedirs(os.path.join(PUBLIC_DIR, "messy"), exist_ok=True)
    os.makedirs(SEALED_DIR, exist_ok=True)
    write_csv(os.path.join(PUBLIC_DIR, "quotes.csv"), CSV_COLS, [csv_row(r) for r in rows])
    write_ledger(np.random.default_rng(SEED + 11), rows, os.path.join(PUBLIC_DIR, "messy", "見積台帳_R5年度_営業1課.xlsx"))
    write_copies(np.random.default_rng(SEED + 12), rows, os.path.join(PUBLIC_DIR, "messy", "見積書控え_2024下期.xlsx"))
    write_memo(np.random.default_rng(SEED + 13), rows, os.path.join(PUBLIC_DIR, "messy", "見積メモ_高橋.xlsx"))
    with open(os.path.join(PUBLIC_DIR, "SCHEMA.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write(SCHEMA_MD.format(nrows=len(rows)))
    with open(os.path.join(SEALED_DIR, "RULES.md"), "w", encoding="utf-8", newline="\n") as f:
        f.write(rules_md(rows))
    write_csv(os.path.join(SEALED_DIR, "truth.csv"), TRUTH_COLS, [truth_row(r) for r in rows])
    lines = []
    for name, p in [("generator.py", os.path.abspath(__file__)), ("RULES.md", os.path.join(SEALED_DIR, "RULES.md")),
                    ("truth.csv", os.path.join(SEALED_DIR, "truth.csv"))]:
        lines.append("%s  %s" % (sha256_file(p), name))
    lines.append("# quotes.csv data rows: %d" % len(rows))
    with open(os.path.join(PUBLIC_DIR, "SEALED.sha256"), "w", encoding="utf-8", newline="\n") as f:
        f.write("\n".join(lines) + "\n")
    print("rows", len(rows))


if __name__ == "__main__":
    main()
