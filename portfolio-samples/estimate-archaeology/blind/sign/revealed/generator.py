#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""SEALED generator for the signage-shop blind-test quote dataset.

Never copy this file (or anything derived from its constants) into the repository.
Deterministic: fixed seeds, no wall-clock values in any output.

Usage:
    python3 generator.py                                   # default locations
    python3 generator.py --repo-out DIR --sealed-out DIR   # alternate locations (self-check)
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import io
import re
import zipfile
import datetime as dt
from pathlib import Path

import numpy as np
import openpyxl
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.utils import get_column_letter

HERE = Path(__file__).resolve().parent
DEFAULT_REPO_OUT = Path(
    "/home/user/hoshi-takarajima-pwa/portfolio-samples/estimate-archaeology/blind/sign")
DEFAULT_SEALED_OUT = HERE

D = dt.date
START, END = D(2022, 10, 1), D(2025, 9, 30)
SEED = 521031
_SS = np.random.SeedSequence(SEED).spawn(5)
RNG_CUST, RNG_Q, RNG_QUIRK, RNG_ORDER, RNG_MESSY = [np.random.default_rng(s) for s in _SS]

# ----------------------------------------------------------------------------
# Calendar
# ----------------------------------------------------------------------------
HOLIDAYS = {D(*x) for x in [
    (2022, 10, 10), (2022, 11, 3), (2022, 11, 23),
    (2023, 1, 2), (2023, 1, 9), (2023, 2, 23), (2023, 3, 21), (2023, 5, 3), (2023, 5, 4),
    (2023, 5, 5), (2023, 7, 17), (2023, 8, 11), (2023, 9, 18), (2023, 10, 9), (2023, 11, 3),
    (2023, 11, 23),
    (2024, 1, 1), (2024, 1, 8), (2024, 2, 12), (2024, 2, 23), (2024, 3, 20), (2024, 4, 29),
    (2024, 5, 3), (2024, 5, 6), (2024, 7, 15), (2024, 8, 12), (2024, 9, 16), (2024, 9, 23),
    (2024, 10, 14), (2024, 11, 4),
    (2025, 1, 1), (2025, 1, 13), (2025, 2, 11), (2025, 2, 24), (2025, 3, 20), (2025, 4, 29),
    (2025, 5, 5), (2025, 5, 6), (2025, 7, 21), (2025, 8, 11), (2025, 9, 15), (2025, 9, 23),
]}


def is_closed(d: dt.date) -> bool:
    if d.weekday() == 6 or d in HOLIDAYS:
        return True
    if (d.month == 12 and d.day >= 29) or (d.month == 1 and d.day <= 4):
        return True
    if d.month == 8 and 13 <= d.day <= 16:
        return True
    return False


def next_open(d: dt.date, k: int) -> dt.date:
    n = 0
    while n < k:
        d = d + dt.timedelta(days=1)
        if not is_closed(d) and d.weekday() != 5:
            n += 1
    return d


# ----------------------------------------------------------------------------
# Reps
# ----------------------------------------------------------------------------
REPS = ["田中", "佐藤", "渡辺", "鈴木", "高橋"]
REP_START = {"田中": START, "佐藤": START, "渡辺": START, "鈴木": D(2023, 4, 3), "高橋": D(2024, 1, 9)}
REP_END = {"田中": END, "佐藤": END, "渡辺": D(2023, 12, 28), "鈴木": END, "高橋": END}


def rep_active(r, d):
    return REP_START[r] <= d <= REP_END[r]


# ----------------------------------------------------------------------------
# Key dates
# ----------------------------------------------------------------------------
R1, R2, R3 = D(2023, 5, 1), D(2024, 4, 1), D(2025, 7, 1)
GRACE_END = D(2024, 5, 31)
INVOICE = D(2023, 10, 1)
CHAIN_RATE_CHANGE = D(2024, 10, 1)
TAKAHASHI_ROUND_CHANGE = D(2024, 7, 1)
WARD_RENAME = D(2024, 1, 1)
TANAKA_OLD_WARD_UNTIL = D(2024, 3, 31)


def table_version(d, grace):
    v = 0 if d < R1 else 1 if d < R2 else 2 if d < R3 else 3
    if v == 2 and grace and d <= GRACE_END:
        v = 1
    return v


def pick(table, v):
    return table[max(k for k in table if k <= v)]


# ----------------------------------------------------------------------------
# Price tables (hidden)
# ----------------------------------------------------------------------------
ACM_SHEET = {0: {"3x6": 5200, "4x8": 9800}, 1: {"3x6": 6400, "4x8": 12000}}
SHEET_DIM = {"3x6": (910, 1820), "4x8": (1220, 2440)}
PRINT_M2 = {0: 3200, 3: 3500}
LAM_M2 = {0: 1200, 3: 1300}
ACM_PROC = {0: 1200, 2: 1500}
ACM_CORNER_R = 400
ACM_HOLES = 300
ACM_JOINT = 6000
OUTPUT_SETUP = 2000
ACM_MIN_TENTHS = 3                       # 0.3 m2 minimum billed per panel face
ACM_TIERS = [(300, 20), (100, 10)]        # billed output area (0.1 m2 units) -> % off output
IJ_MEDIA = {0: {"塩ビ": 2800, "ターポリン": 3000, "合成紙": 2200, "透過フィルム": 4200},
            3: {"塩ビ": 3100, "ターポリン": 3300, "合成紙": 2400, "透過フィルム": 4600}}
IJ_MIN_TENTHS = 2
IJ_ROLL_SHORT_LIMIT = 900
IJ_ROLL_WIDTH_MILLI = 1370                # 1.37 m in mm
IJ_ROLL_PRINTABLE = 1340
IJ_TIERS = [(50000, 20), (20000, 10)]     # billed m2*1000 -> % off media+lam
GROMMET = {0: 60, 3: 70}
GROMMET_PITCH = 500
CALP_BANDS = {0: [(100, 1200), (150, 1600), (200, 2100), (300, 3000), (400, 4200), (500, 5500), (600, 7000)],
              2: [(100, 1300), (150, 1750), (200, 2300), (300, 3300), (400, 4600), (500, 6000), (600, 7600)]}
CALP_OVER = {0: 1500, 2: 1650}
CALP_THICK = {10: 85, 20: 100, 30: 120, 50: 150}
CALP_FINISH = {"シート貼り": 100, "塗装": 135}
CALP_FULL_UPTO = 20
CALP_EXCESS_PCT = 85
CALP_SETUP = {0: 3000, 2: 3500}
SODE_CLASSES = {0: [((450, 900), 85000), ((600, 1200), 118000), ((600, 1800), 152000),
                    ((750, 1800), 178000), ((900, 2400), 236000)],
                1: [((450, 900), 95000), ((600, 1200), 132000), ((600, 1800), 170000),
                    ((750, 1800), 199000), ((900, 2400), 264000)]}
SODE_OVER = {0: 28000, 1: 31000}
SODE_LED = {0: (38000, 15000), 1: (42000, 16000)}
MIN_PRODUCT = {0: 5000, 2: 6000}
RUSH_PCT, RUSH_MIN = 30, 3000
DESIGN = {0: {"修正": 3000, "新規": {"ACM": 10000, "CALP": 6000, "IJ": 8000, "SODE": 15000}},
          3: {"修正": 3500, "新規": {"ACM": 12000, "CALP": 7000, "IJ": 9000, "SODE": 18000}}}
INSTALL_ACM = {0: (5000, 3000), 2: (6000, 3500)}
INSTALL_CALP = {0: 700, 2: 800}
INSTALL_IJ = {0: 1800, 2: 2000}
INSTALL_SODE = {0: 38000, 2: 45000}
INSTALL_MIN = {0: 12000, 2: 15000}
NIGHT_PCT = 50
ELECTRICAL = {0: 18000, 2: 22000}
HEIGHT = {0: [(20, 0), (40, 8000), (90, 28000), (None, 48000)],
          2: [(20, 0), (40, 10000), (90, 33000), (None, 55000)]}
REMOVAL = {0: 15000, 2: 18000}
DISPOSAL = 5000
SURVEY = {0: 5000, 2: 6000}
TRAVEL = {0: ([(10, 0), (25, 4000), (50, 9000), (80, 16000)], 2000),
          2: ([(10, 0), (25, 5000), (50, 12000), (80, 20000)], 2500)}

ITEM_NAME = {"ACM": "アルミ複合板看板", "CALP": "カルプ文字", "IJ": "インクジェット出力", "SODE": "袖看板"}
ITEM_UNIT = {"ACM": "枚", "CALP": "文字", "IJ": "枚", "SODE": "台"}

# Sites: (name before ward reorganisation, name from 2024-01-01, km from shop, weight)
SITES = [
    ("浜松市中区鍛冶町", "浜松市中央区鍛冶町", 1, 10),
    ("浜松市中区佐鳴台", "浜松市中央区佐鳴台", 5, 6),
    ("浜松市東区天王町", "浜松市中央区天王町", 6, 6),
    ("浜松市南区高塚町", "浜松市中央区高塚町", 8, 5),
    ("浜松市北区三方原町", "浜松市中央区三方原町", 10, 4),
    ("浜松市西区雄踏町", "浜松市中央区雄踏町", 11, 4),
    ("浜松市浜北区貴布祢", "浜松市浜名区貴布祢", 13, 5),
    ("浜松市西区舞阪町", "浜松市中央区舞阪町", 14, 3),
    ("磐田市", "磐田市", 18, 6),
    ("浜松市北区細江町", "浜松市浜名区細江町", 19, 3),
    ("湖西市", "湖西市", 22, 4),
    ("浜松市北区引佐町", "浜松市浜名区引佐町", 24, 2),
    ("浜松市天竜区二俣町", "浜松市天竜区二俣町", 26, 2),
    ("袋井市", "袋井市", 27, 4),
    ("周智郡森町", "周智郡森町", 33, 2),
    ("豊橋市", "豊橋市", 35, 4),
    ("掛川市", "掛川市", 38, 3),
    ("菊川市", "菊川市", 42, 2),
    ("豊川市", "豊川市", 48, 2),
    ("新城市", "新城市", 52, 1),
    ("御前崎市", "御前崎市", 55, 1),
    ("島田市", "島田市", 58, 1),
    ("藤枝市", "藤枝市", 66, 1),
    ("岡崎市", "岡崎市", 75, 1),
    ("静岡市駿河区", "静岡市駿河区", 80, 1),
    ("静岡市葵区", "静岡市葵区", 84, 1),
    ("名古屋市中区", "名古屋市中区", 108, 0.5),
]


def site_name(idx, d, rep):
    old, new, _, _ = SITES[idx]
    if d < WARD_RENAME:
        return old
    if rep == "田中" and d <= TANAKA_OLD_WARD_UNTIL:
        return old
    return new


# ----------------------------------------------------------------------------
# Remarks vocabulary
# ----------------------------------------------------------------------------
RUSH_P = ["特急", "至急", "特急対応", "至急でお願いします", "【特急】", "特急（3営業日）", "至急対応希望"]
NIGHT_P = ["夜間作業", "夜間施工", "閉店後夜間施工", "夜間（22時以降）"]
REMOVAL_ONLY_P = ["既存看板撤去", "既存撤去あり", "既存サイン撤去（処分は先方）", "既存撤去（処分不要）"]
REMOVAL_DISP_P = ["既存看板撤去・処分", "撤去処分込み", "既存撤去・産廃処分", "既存看板撤去／処分"]
SURVEY_P = ["現調あり", "現地調査希望", "要現調", "現地調査含む"]
CORNER_P = ["角R加工", "角丸", "四隅R"]
HOLES_P = ["穴あけ", "ビス穴加工", "穴あけ4箇所"]
NOGROM_P = ["ハトメなし", "ハトメ不要"]
REPEAT_P = ["前回同様", "リピート", "増刷", "前回データで"]
DECOY_RUSH = ["納期急ぎません", "急ぎではない", "なるべく急ぎで", "お急ぎとのこと", "早めの対応希望"]
DECOY_SURVEY = ["現調不要", "現地調査不要（写真あり）"]
DECOY_REMOVAL = ["撤去は別途"]
NOISE_GENERAL = ["色校正あり", "PDF送付済", "見積有効期限1ヶ月", "電話にて依頼", "メールにて依頼",
                 "ロゴデータAI支給", "データ入稿済", "2案提出", "支払：月末締翌月末", "ご担当：店長様",
                 "オープン前", "新店舗", "改装に伴う", "テナント看板", "仮見積", "概算", "社内確認中",
                 "白フチあり", "背景：白", "フルカラー", "2色", "ロゴのみ", "電話番号変更", "店名変更",
                 "期間限定キャンペーン用", "イベント用", "移転に伴い", "引取り", "配送希望", "校正2回まで",
                 "屋外用", "屋内用", "紹介案件", "FAXにて送付", "請求書は本社宛"]
NOISE_INSTALL = ["週末施工希望", "駐車場あり", "雨天順延", "現場写真受領", "前面道路狭い", "搬入経路要確認",
                 "管理会社承認待ち", "要図面", "午前中施工希望", "平日のみ施工可"]


def parse_flags(t: str) -> dict:
    f = {}
    f["rush"] = ("特急" in t) or ("至急" in t)
    f["night"] = "夜間" in t
    f["removal"] = ("撤去" in t) and ("撤去は別途" not in t)
    f["disposal"] = f["removal"] and ("処分" in t) and not any(
        x in t for x in ("処分は先方", "処分不要", "処分なし"))
    f["survey"] = (("現調" in t) or ("現地調査" in t)) and not any(
        x in t for x in ("現調不要", "現地調査不要"))
    f["corner_r"] = any(x in t for x in ("角R", "角丸", "四隅R"))
    f["holes"] = any(x in t for x in ("穴あけ", "ビス穴"))
    f["no_grommet"] = any(x in t for x in ("ハトメなし", "ハトメ不要"))
    f["repeat"] = any(x in t for x in ("前回同様", "リピート", "増刷", "前回データ"))
    return f


FLAG_KEYS = ["rush", "night", "removal", "disposal", "survey", "corner_r", "holes", "no_grommet", "repeat"]


# ----------------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------------
def cdiv(a, b):
    return -(-a // b)


def choice(rng, items, weights=None):
    if weights is None:
        return items[int(rng.integers(len(items)))]
    w = np.asarray(weights, dtype=float)
    return items[int(rng.choice(len(items), p=w / w.sum()))]


def per_sheet(W, H, SW, SH):
    return max((SW // W) * (SH // H), (SW // H) * (SH // W))


# ----------------------------------------------------------------------------
# Customers
# ----------------------------------------------------------------------------
TRADE_CODES = [1005, 1012, 1019, 1030, 1036, 1047, 1058, 1063, 1071, 1084]
SPECIAL = {1023: "chain", 1041: "contractor", 1077: "regular", 1088: "general"}

PRODUCT_MIX = {
    "general": ([0.36, 0.30, 0.22, 0.12]),
    "regular": ([0.38, 0.30, 0.18, 0.14]),
    "trade": ([0.30, 0.62, 0.08, 0.00]),
    "chain": ([0.45, 0.40, 0.05, 0.10]),
    "contractor": ([0.40, 0.05, 0.35, 0.20]),
    "walkin": ([0.38, 0.35, 0.18, 0.09]),
}
PRODUCTS = ["ACM", "IJ", "CALP", "SODE"]
INSTALL_FACTOR = {"general": 1.0, "regular": 1.0, "trade": 0.15, "chain": 1.45, "contractor": 1.5, "walkin": 0.9}
DESIGN_MIX = {"general": [0.35, 0.25, 0.40], "regular": [0.50, 0.25, 0.25], "trade": [0.90, 0.08, 0.02],
              "chain": [0.50, 0.35, 0.15], "contractor": [0.45, 0.30, 0.25], "walkin": [0.30, 0.25, 0.45]}


def build_customers(rng):
    pool = [c for c in range(1001, 1097) if c not in TRADE_CODES and c not in SPECIAL]
    drop = set(int(x) for x in rng.choice(pool, size=14, replace=False))
    codes = sorted(c for c in range(1001, 1097) if c not in drop)
    others = [c for c in codes if c not in TRADE_CODES and c not in SPECIAL]
    regular = set(int(x) for x in rng.choice(others, size=20, replace=False))
    custs = []
    for c in codes:
        if c in SPECIAL:
            cls = SPECIAL[c]
        elif c in TRADE_CODES:
            cls = "trade"
        elif c in regular:
            cls = "regular"
        else:
            cls = "general"
        u = rng.random()
        if u < 0.68:
            start, end = START, END
        elif u < 0.88:
            start = START + dt.timedelta(days=int(rng.integers(30, 900)))
            end = END
        else:
            start = START
            end = START + dt.timedelta(days=int(rng.integers(200, 1000)))
        w = float(rng.lognormal(0.0, 0.85))
        if cls == "regular":
            w *= 1.6
        if cls == "trade":
            w *= 1.3
        if start < D(2023, 4, 1):
            rep = choice(rng, ["田中", "佐藤", "渡辺"], [0.25, 0.45, 0.30])
        elif start < D(2024, 1, 1):
            rep = choice(rng, ["田中", "佐藤", "鈴木", "渡辺"], [0.15, 0.40, 0.40, 0.05])
        else:
            rep = choice(rng, ["佐藤", "鈴木", "高橋"], [0.25, 0.35, 0.40])
        if rep == "渡辺" and start > D(2023, 9, 30):
            rep = "佐藤"
        cust = dict(code=str(c), cls=cls, start=start, end=end, w=w, rep=rep,
                    travel_waived=(c == 1077), basis="incl" if c == 1088 else "excl",
                    grace=cls in ("regular", "chain", "contractor"))
        if c == 1023:
            cust.update(w=9.0, rep="田中", start=START, end=END)
        if c == 1041:
            cust.update(w=3.2, start=START, end=END)
        if c == 1077:
            cust.update(w=2.4, start=START, end=END)
        if c == 1088:
            cust.update(w=1.4)
        custs.append(cust)
    custs.append(dict(code="9999", cls="walkin", start=START, end=END, w=30.0, rep=None,
                      travel_waived=False, basis=None, grace=False))
    return custs


def cust_mult(cust, d):
    if cust["code"] == "1023":
        return 880 if d < CHAIN_RATE_CHANGE else 850
    return {"general": 1000, "regular": 950, "trade": 800, "contractor": 900, "walkin": 1000}[cust["cls"]]


def rep_for(cust, d, rng, saturday):
    if saturday:
        return "田中"
    active = [r for r in REPS if rep_active(r, d)]
    p = cust["rep"]
    if p is None or rng.random() < 0.12:
        weights = {"田中": 0.8, "佐藤": 1.2, "渡辺": 1.0, "鈴木": 1.0, "高橋": 1.0}
        return choice(rng, active, [weights[r] for r in active])
    if p == "渡辺" and d > REP_END["渡辺"]:
        p = "高橋" if d >= REP_START["高橋"] else "佐藤"
    if not rep_active(p, d):
        p = "佐藤"
    return p


# ----------------------------------------------------------------------------
# Rounding / tax
# ----------------------------------------------------------------------------
def rep_round(rep, d, x):
    if rep == "田中":
        return (x // 1000 * 1000, "floor1000") if x >= 30000 else (x // 100 * 100, "floor100")
    if rep == "渡辺":
        return ((x + 50) // 100 * 100, "halfup100")
    if rep == "高橋" and d < TAKAHASHI_ROUND_CHANGE:
        return (x // 10 * 10, "floor10")
    return (x // 100 * 100, "floor100")


def tax_label_for(rep, d, basis):
    if basis == "incl":
        return "税込"
    if d < INVOICE:
        return {"田中": "税別", "佐藤": "税別", "渡辺": "", "鈴木": "税抜", "高橋": "税抜"}[rep]
    return "" if rep == "渡辺" else "税抜"


def written_from_net(net, d, basis):
    if basis == "excl":
        return net
    if d < INVOICE:
        return (net * 110 // 100) // 100 * 100
    return net + net * 10 // 100


# ----------------------------------------------------------------------------
# Pricing
# ----------------------------------------------------------------------------
COMP_KEYS = ["material", "output_gross", "volume_discount", "output", "processing", "setup",
             "product_raw", "mult", "product_after_mult", "min_charge_adj", "product_final", "rush",
             "design_fee", "install_base_raw", "install_min_adj", "install_base", "night_extra",
             "height_fee", "electrical", "removal", "disposal", "survey", "travel_km", "travel",
             "pre_round_total"]


def price(q, *, v_shift=0, mult_override=None, force_rush=False, omit_travel=False):
    d, f, it = q["date"], q["flags"], q["item"]
    v = table_version(d, q["grace"]) + v_shift
    c = {k: 0 for k in COMP_KEYS}
    c["table_version"] = v
    W, H, n = q["W"], q["H"], q["qty"]
    if it == "ACM":
        if per_sheet(W, H, *SHEET_DIM["3x6"]) >= 1:
            c["material"] = cdiv(n, per_sheet(W, H, *SHEET_DIM["3x6"])) * pick(ACM_SHEET, v)["3x6"]
        elif per_sheet(W, H, *SHEET_DIM["4x8"]) >= 1:
            c["material"] = cdiv(n, per_sheet(W, H, *SHEET_DIM["4x8"])) * pick(ACM_SHEET, v)["4x8"]
        else:
            tiles = min(cdiv(W, 1220) * cdiv(H, 2440), cdiv(H, 1220) * cdiv(W, 2440))
            c["material"] = n * tiles * pick(ACM_SHEET, v)["4x8"]
            c["processing"] += n * (tiles - 1) * ACM_JOINT
        area_t = max(ACM_MIN_TENTHS, cdiv(W * H, 100000))
        sides = 2 if q["sides"] == "両面" else 1
        out_t = area_t * n * sides
        unit = pick(PRINT_M2, v) + (pick(LAM_M2, v) if q["lam"] != "なし" else 0)
        c["output_gross"] = unit * out_t // 10
        pct = next((p for lim, p in ACM_TIERS if out_t >= lim), 0)
        c["volume_discount"] = c["output_gross"] * pct // 100
        c["output"] = c["output_gross"] - c["volume_discount"]
        c["processing"] += n * pick(ACM_PROC, v)
        if f["corner_r"]:
            c["processing"] += n * ACM_CORNER_R
        if f["holes"]:
            c["processing"] += n * ACM_HOLES
        c["setup"] = 0 if f["repeat"] else OUTPUT_SETUP
    elif it == "IJ":
        s, l = min(W, H), max(W, H)
        if s <= IJ_ROLL_SHORT_LIMIT:
            milli = max(IJ_MIN_TENTHS, cdiv(W * H, 100000)) * 100
        else:
            strips = cdiv(s, IJ_ROLL_PRINTABLE)
            milli = strips * IJ_ROLL_WIDTH_MILLI * cdiv(l, 100) // 10
        tot = milli * n
        media = pick(IJ_MEDIA, v)[q["media"]] * tot // 1000
        lam = (pick(LAM_M2, v) * tot // 1000) if q["lam"] != "なし" else 0
        c["output_gross"] = media + lam
        pct = next((p for lim, p in IJ_TIERS if tot >= lim), 0)
        c["volume_discount"] = c["output_gross"] * pct // 100
        c["output"] = c["output_gross"] - c["volume_discount"]
        if q["media"] == "ターポリン" and not f["no_grommet"]:
            g = 2 * (cdiv(W, GROMMET_PITCH) + cdiv(H, GROMMET_PITCH)) * n
            c["processing"] = g * pick(GROMMET, v)
        c["setup"] = 0 if f["repeat"] else OUTPUT_SETUP
    elif it == "CALP":
        h = H
        bands = pick(CALP_BANDS, v)
        base = next((p for lim, p in bands if h <= lim), None)
        if base is None:
            base = bands[-1][1] + cdiv(h - 600, 100) * pick(CALP_OVER, v)
        raw = base * CALP_THICK[q["thick"]] * CALP_FINISH[q["finish"]]
        unit = cdiv(raw, 100000) * 10
        full = min(n, CALP_FULL_UPTO)
        extra = max(0, n - CALP_FULL_UPTO)
        c["output_gross"] = unit * n
        c["output"] = unit * full + (unit * CALP_EXCESS_PCT // 100) * extra
        c["volume_discount"] = c["output_gross"] - c["output"]
        c["setup"] = 0 if f["repeat"] else pick(CALP_SETUP, v)
    elif it == "SODE":
        s, l = min(W, H), max(W, H)
        classes = pick(SODE_CLASSES, v)
        base = next((p for (cs, cl), p in classes if s <= cs and l <= cl), None)
        if base is None:
            base = classes[-1][1] + cdiv(W * H - 2160000, 500000) * pick(SODE_OVER, v)
        c["material"] = base * n
        if q["light"] == "LED内照":
            fixed, per = pick(SODE_LED, v)
            c["processing"] = (fixed + per * cdiv(W * H, 100000) // 10) * n
    c["product_raw"] = c["material"] + c["output"] + c["processing"] + c["setup"]
    m = mult_override if mult_override is not None else cust_mult(q["cust"], d)
    c["mult"] = m
    c["product_after_mult"] = c["product_raw"] * m // 1000
    mn = pick(MIN_PRODUCT, v)
    c["product_final"] = max(mn, c["product_after_mult"])
    c["min_charge_adj"] = c["product_final"] - c["product_after_mult"]
    if f["rush"] or force_rush:
        c["rush"] = max(RUSH_MIN, c["product_final"] * RUSH_PCT // 100)
    if q["design"] != "支給":
        dz = pick(DESIGN, v)
        c["design_fee"] = dz["修正"] if q["design"] == "修正" else dz["新規"][it]
    if q["install"]:
        if it == "ACM":
            b, per = pick(INSTALL_ACM, v)
            raw = b + per * cdiv(W * H * n, 100000) // 10
        elif it == "CALP":
            raw = pick(INSTALL_CALP, v) * n
        elif it == "IJ":
            raw = pick(INSTALL_IJ, v) * cdiv(W * H * n, 100000) // 10
        else:
            raw = pick(INSTALL_SODE, v) * n
        c["install_base_raw"] = raw
        c["install_base"] = max(pick(INSTALL_MIN, v), raw)
        c["install_min_adj"] = c["install_base"] - raw
        if f["night"]:
            c["night_extra"] = c["install_base"] * NIGHT_PCT // 100
        ht = q["height_t"]
        c["height_fee"] = next(fee for lim, fee in pick(HEIGHT, v) if lim is None or ht <= lim)
        if it == "SODE" and q["light"] == "LED内照":
            c["electrical"] = pick(ELECTRICAL, v)
        if f["removal"]:
            c["removal"] = pick(REMOVAL, v)
            if f["disposal"]:
                c["disposal"] = DISPOSAL
    if f["survey"]:
        c["survey"] = pick(SURVEY, v)
    if (q["install"] or f["survey"]) and q["site"] is not None:
        km = SITES[q["site"]][2]
        c["travel_km"] = km
        if not q["cust"]["travel_waived"] and not omit_travel:
            bands, step = pick(TRAVEL, v)
            fee = next((fee for lim, fee in bands if km <= lim), None)
            if fee is None:
                fee = bands[-1][1] + cdiv(km - 80, 10) * step
            c["travel"] = fee
    c["pre_round_total"] = (c["product_final"] + c["rush"] + c["design_fee"] + c["install_base"]
                            + c["night_extra"] + c["height_fee"] + c["electrical"] + c["removal"]
                            + c["disposal"] + c["survey"] + c["travel"])
    net, rule = rep_round(q["rep"], d, c["pre_round_total"])
    c["rounding_rule"] = rule
    c["rounding_adj"] = net - c["pre_round_total"]
    c["net"] = net
    return c


# ----------------------------------------------------------------------------
# Spec generation
# ----------------------------------------------------------------------------
ACM_SIZES = [((900, 1800), 12), ((600, 900), 10), ((450, 600), 8), ((600, 1800), 6), ((1200, 2400), 4),
             ((900, 900), 5), ((300, 450), 6), ((1800, 900), 5), ((1500, 900), 4), ((910, 1820), 4),
             ((450, 900), 4), ((600, 600), 4), ((1000, 2000), 2), ((1820, 3640), 1), ((2400, 1200), 2),
             ((3000, 900), 1), ((1200, 600), 3), ((750, 1500), 2)]
IJ_SIZES = [((594, 841), 8), ((841, 1189), 6), ((728, 1030), 5), ((1030, 1456), 2), ((420, 594), 6),
            ((900, 1800), 6), ((1200, 1800), 4), ((900, 3600), 4), ((1800, 900), 4), ((450, 1200), 3),
            ((600, 1800), 4), ((1000, 3000), 3), ((1370, 5000), 1), ((300, 900), 3), ((515, 728), 3)]
SODE_SIZES = [((450, 900), 10), ((450, 1200), 8), ((600, 1200), 10), ((600, 1800), 8), ((700, 1800), 5),
              ((750, 1500), 3), ((900, 2400), 4), ((1000, 3000), 2), ((500, 1500), 4), ((400, 800), 3)]
CALP_HEIGHTS = [80, 100, 120, 150, 180, 200, 250, 300, 350, 400, 450, 500, 600, 700, 800, 900]
CALP_HW = [3, 6, 5, 9, 4, 10, 8, 10, 5, 7, 3, 4, 3, 2, 1, 1]
H_ACM = [15, 18, 20, 22, 25, 30, 35, 40, 45, 50, 60, 80]
H_ACM_W = [6, 8, 10, 6, 9, 12, 9, 8, 5, 4, 3, 1]
H_CALP = [20, 25, 30, 35, 40, 45, 50, 60, 75]
H_CALP_W = [5, 8, 12, 10, 8, 6, 5, 3, 1]
H_IJ = [10, 15, 18, 20, 24, 30, 35]
H_IJ_W = [6, 10, 10, 10, 5, 4, 2]
H_SODE = [30, 35, 40, 45, 50, 60, 70, 80, 90, 100, 120]
H_SODE_W = [4, 6, 8, 8, 9, 9, 7, 6, 5, 4, 2]


def draw_size(rng, table, rand_p, lo, hi):
    if rng.random() < rand_p:
        return int(rng.integers(lo // 50, hi // 50 + 1)) * 50, int(rng.integers(lo // 50, hi // 50 + 1)) * 50
    sizes, ws = zip(*table)
    return choice(rng, list(sizes), ws)


def draw_qty(rng, profile):
    u = rng.random()
    if profile == "acm":
        if u < .55: return 1
        if u < .75: return 2
        if u < .88: return int(rng.integers(3, 6))
        if u < .97: return int(rng.integers(6, 21))
        return int(rng.integers(21, 61))
    if profile == "ij":
        if u < .35: return 1
        if u < .65: return int(rng.integers(2, 6))
        if u < .87: return int(rng.integers(6, 21))
        return int(rng.integers(21, 101))
    if profile == "calp":
        return int(min(60, max(1, round(rng.lognormal(2.2, 0.6)))))
    if profile == "sode":
        return 1 if u < .9 else 2
    raise ValueError(profile)


def gen_quote(rng, d, cust, rep):
    cls = cust["cls"]
    it = choice(rng, PRODUCTS, PRODUCT_MIX[cls])
    q = dict(date=d, cust=cust, rep=rep, item=it, W=None, H=None, qty=None, thick=None, sides="",
             media="", lam="", finish="", light="", design="支給", install=False, height_t=None,
             site=None, grace=cust["grace"])
    inst_p = {"ACM": .55, "CALP": .70, "IJ": .30, "SODE": .90}[it] * INSTALL_FACTOR[cls]
    if it == "ACM":
        if rng.random() < 0.05:
            q["W"], q["H"] = int(rng.integers(30, 73)) * 50, int(rng.integers(18, 49)) * 50
        else:
            q["W"], q["H"] = draw_size(rng, ACM_SIZES, 0.30, 200, 1950)
        q["qty"] = draw_qty(rng, "acm")
        q["thick"] = 3
        q["sides"] = "両面" if rng.random() < .15 else "片面"
        q["lam"] = choice(rng, ["グロス", "マット", "なし"], [.55, .35, .10])
        q["qty"] = min(q["qty"], max(1, 30_000_000 // (q["W"] * q["H"])))
    elif it == "IJ":
        q["W"], q["H"] = draw_size(rng, IJ_SIZES, 0.30, 200, 1800)
        q["qty"] = draw_qty(rng, "ij")
        if cls == "trade" and rng.random() < .4:
            q["qty"] = int(rng.integers(10, 120))
        q["media"] = choice(rng, ["塩ビ", "ターポリン", "合成紙", "透過フィルム"], [.45, .25, .20, .10])
        q["lam"] = {"塩ビ": lambda: choice(rng, ["グロス", "マット", "なし"], [.45, .35, .20]),
                    "合成紙": lambda: choice(rng, ["グロス", "マット", "なし"], [.3, .3, .4]),
                    "透過フィルム": lambda: choice(rng, ["グロス", "なし"], [.5, .5]),
                    "ターポリン": lambda: choice(rng, ["なし", "マット"], [.9, .1])}[q["media"]]()
        q["qty"] = min(q["qty"], max(1, 80_000_000 // (q["W"] * q["H"])))
    elif it == "CALP":
        q["H"] = choice(rng, CALP_HEIGHTS, CALP_HW)
        q["qty"] = draw_qty(rng, "calp")
        if q["H"] >= 600:
            q["qty"] = min(q["qty"], 12)
        q["thick"] = choice(rng, [10, 20, 30, 50], [.15, .45, .30, .10])
        q["finish"] = choice(rng, ["シート貼り", "塗装"], [.65, .35])
    else:
        q["W"], q["H"] = draw_size(rng, SODE_SIZES, 0.0, 0, 0)
        q["qty"] = draw_qty(rng, "sode")
        q["sides"] = "両面"
        q["light"] = choice(rng, ["LED内照", "なし"], [.55, .45])
    if cls == "walkin":
        q["qty"] = min(q["qty"], 10)
    q["unit"] = ITEM_UNIT[it]
    q["install"] = bool(rng.random() < min(0.95, inst_p))
    q["design"] = choice(rng, ["支給", "修正", "新規"], DESIGN_MIX[cls])
    flags = {k: False for k in FLAG_KEYS}
    if cls != "walkin" and rng.random() < (0.15 if cls == "trade" else 0.09):
        flags["repeat"] = True
        q["design"] = "支給"
    flags["rush"] = bool(rng.random() < (0.08 if cls == "trade" else 0.05))
    if q["install"]:
        hl, hw = {"ACM": (H_ACM, H_ACM_W), "CALP": (H_CALP, H_CALP_W), "IJ": (H_IJ, H_IJ_W),
                  "SODE": (H_SODE, H_SODE_W)}[it]
        q["height_t"] = choice(rng, hl, hw)
        flags["night"] = bool(rng.random() < (0.15 if cls == "chain" else 0.05))
        if rng.random() < (0.20 if cls == "chain" else 0.13):
            flags["removal"] = True
            flags["disposal"] = bool(rng.random() < 0.65)
        big = it == "SODE" or (q["W"] or 0) * (q["H"] or 0) * q["qty"] >= 2_000_000
        flags["survey"] = bool(rng.random() < (0.20 if big else 0.06))
    else:
        flags["survey"] = bool(rng.random() < 0.02)
    if it == "ACM":
        flags["corner_r"] = bool(rng.random() < .10)
        flags["holes"] = bool(rng.random() < .12)
    if it == "IJ" and q["media"] == "ターポリン":
        flags["no_grommet"] = bool(rng.random() < .12)
    if q["install"] or flags["survey"]:
        sw = [s[3] for s in SITES]
        if cls == "chain":
            sw = [max(1.0, w) for w in sw]
        elif cls == "walkin":
            sw = [w if s[2] <= 30 else w * 0.3 for w, s in zip(sw, SITES)]
        q["site"] = choice(rng, list(range(len(SITES))), sw)
    q["planned_flags"] = flags
    q["remarks"] = build_remarks(rng, q, flags)
    q["flags"] = parse_flags(q["remarks"])
    assert q["flags"] == flags, (q["remarks"], flags, q["flags"])
    return q


def store_name(idx):
    nm = SITES[idx][1]
    if nm.startswith("浜松市"):
        nm = nm.split("区", 1)[1]
        nm = nm[:-1] if nm.endswith("町") else nm
    elif nm.endswith("区"):
        city, ward = nm.split("市", 1)
        nm = city + ward[:-1]
    else:
        nm = nm.replace("周智郡", "")
        nm = nm[:-1] if nm[-1] in "市町" else nm
    return f"{nm}店"


def build_remarks(rng, q, flags):
    parts = []
    if flags["repeat"]:
        parts.append(choice(rng, REPEAT_P))
    if flags["rush"]:
        parts.append(choice(rng, RUSH_P))
    elif rng.random() < 0.035:
        parts.append(choice(rng, DECOY_RUSH))
    if flags["night"]:
        parts.append(choice(rng, NIGHT_P))
    if flags["removal"]:
        parts.append(choice(rng, REMOVAL_DISP_P if flags["disposal"] else REMOVAL_ONLY_P))
    elif q["install"] and rng.random() < 0.025:
        parts.append(choice(rng, DECOY_REMOVAL))
    if flags["survey"]:
        parts.append(choice(rng, SURVEY_P))
    elif q["install"] and rng.random() < 0.02:
        parts.append(choice(rng, DECOY_SURVEY))
    if flags["corner_r"]:
        parts.append(choice(rng, CORNER_P))
    if flags["holes"]:
        parts.append(choice(rng, HOLES_P))
    if flags["no_grommet"]:
        parts.append(choice(rng, NOGROM_P))
    k = int(rng.choice([0, 1, 2, 3], p=[.38, .37, .18, .07]))
    pool = NOISE_GENERAL + (NOISE_INSTALL if q["install"] else [])
    for i in rng.choice(len(pool), size=k, replace=False):
        parts.append(pool[int(i)])
    if q["cust"]["cls"] == "chain" and q["site"] is not None and rng.random() < 0.6:
        parts.append(store_name(q["site"]))
    order = rng.permutation(len(parts))
    parts = [parts[int(i)] for i in order]
    sep = choice(rng, ["、", "／", " "], [.6, .2, .2])
    return sep.join(parts)


# ----------------------------------------------------------------------------
# Base quote generation
# ----------------------------------------------------------------------------
SEASON = {1: .78, 2: 1.0, 3: 1.32, 4: 1.15, 5: .95, 6: 1.0, 7: 1.02, 8: .80, 9: 1.05, 10: 1.10, 11: 1.0, 12: .90}
BASE_LAMBDA = 2.95


def generate_base(custs, rng):
    quotes = []
    d = START
    while d <= END:
        if not is_closed(d):
            sat = d.weekday() == 5
            trend = 1.0 + 0.07 * ((d - START).days / 365.0)
            lam = BASE_LAMBDA * SEASON[d.month] * trend * (0.12 if sat else 1.0)
            n = int(rng.poisson(lam))
            active = [c for c in custs if c["start"] <= d <= c["end"]]
            ws = [c["w"] for c in active]
            for _ in range(n):
                cust = choice(rng, active, ws)
                rep = rep_for(cust, d, rng, sat)
                q = gen_quote(rng, d, cust, rep)
                if cust["basis"] is None:
                    q["basis"] = "incl" if rng.random() < 0.6 else "excl"
                else:
                    q["basis"] = cust["basis"]
                quotes.append(q)
        d += dt.timedelta(days=1)
    return quotes


# ----------------------------------------------------------------------------
# Quirks
# ----------------------------------------------------------------------------
def digit_typo(n, rng):
    s = str(n)
    for _ in range(50):
        kind = choice(rng, ["swap", "sub", "drop0", "add0"], [.4, .35, .15, .10])
        t = s
        if kind == "swap" and len(s) >= 3:
            i = int(rng.integers(0, len(s) - 1))
            if s[i] != s[i + 1] and not (i == 0 and s[1] == "0"):
                t = s[:i] + s[i + 1] + s[i] + s[i + 2:]
        elif kind == "sub":
            sim = {"1": "7", "7": "1", "3": "8", "8": "3", "5": "6", "6": "5", "0": "9", "9": "0",
                   "2": "3", "4": "9"}
            i = int(rng.integers(0, len(s)))
            r = sim[s[i]]
            if not (i == 0 and r == "0"):
                t = s[:i] + r + s[i + 1:]
        elif kind == "drop0" and "0" in s[1:]:
            i = s.index("0", 1)
            t = s[:i] + s[i + 1:]
        elif kind == "add0":
            i = len(s) - 2
            t = s[:i] + "0" + s[i:]
        if t != s:
            return int(t), kind
    raise RuntimeError("typo failed")


def finalize_clean(q):
    q["true_net"] = q["actual"]["net"]
    q["written"] = written_from_net(q["true_net"], q["date"], q["basis"])
    q["tax_label"] = tax_label_for(q["rep"], q["date"], q["basis"])


def pick_apply(rng, cands, weight_fn, k, fn):
    """Weighted random order over candidates; apply fn until k succeed (fn returns True when applied)."""
    cands = list(cands)
    done = []
    while cands and len(done) < k:
        w = np.array([weight_fn(c) for c in cands], dtype=float)
        i = cands.pop(int(rng.choice(len(cands), p=w / w.sum())))
        if fn(i):
            done.append(i)
    return done


def apply_quirks(quotes, rng):
    for q in quotes:
        q["rule"] = price(q)
        q["actual"] = q["rule"]
        q["quirk"], q["detail"], q["post_adj"] = "clean", "", 0
        finalize_clean(q)
    idx = list(range(len(quotes)))
    free = lambda i: quotes[i]["quirk"] == "clean"
    one = lambda i: 1.0

    def set_actual(q, actual):
        q["actual"] = actual
        q["true_net"] = actual["net"]
        q["written"] = written_from_net(q["true_net"], q["date"], q["basis"])

    def set_net(q, v):
        q["true_net"] = v
        q["post_adj"] = v - q["actual"]["net"]
        q["written"] = written_from_net(v, q["date"], q["basis"])

    # --- impossible to reconcile
    modes = iter(["other_quote", "round_unrelated", "bundle", "other_quote", "partial", "round_unrelated"])

    def f_imp(i):
        q = quotes[i]
        r = q["rule"]["net"]
        mode = next(modes)
        if mode == "other_quote":
            others = [j for j in idx if quotes[j]["item"] != q["item"] and abs(quotes[j]["rule"]["net"] - r) > r * 0.6]
            v = quotes[others[int(rng.integers(len(others)))]]["rule"]["net"]
        elif mode == "round_unrelated":
            v = int(choice(rng, [30000, 50000, 80000, 100000, 150000, 200000]))
            if abs(v - r) < r * 0.4:
                v *= 3
        elif mode == "bundle":
            v = int(r * rng.uniform(2.2, 3.5)) // 1000 * 1000
        else:
            v = int(r * rng.uniform(0.30, 0.45)) // 100 * 100
        set_net(q, v)
        q["quirk"], q["detail"] = "impossible", f"mode={mode}; rule_net={r}; written amount implies net {v}"
        return True
    pick_apply(rng, [i for i in idx if quotes[i]["rule"]["net"] >= 8000], one, 6, f_imp)

    # --- stale price table after a revision
    rep_w_stale = {"田中": 2.0, "渡辺": 3.0, "佐藤": 1.0, "鈴木": 1.0, "高橋": 1.2}

    def f_stale(i):
        q = quotes[i]
        alt = price(q, v_shift=-1)
        if alt["net"] == q["rule"]["net"]:
            return False
        set_actual(q, alt)
        q["quirk"] = "stale_table"
        q["detail"] = (f"priced with table v{alt['table_version']} instead of v{q['rule']['table_version']}; "
                       f"diff={q['true_net'] - q['rule']['net']}")
        return True
    windows = [(R1, D(2023, 7, 14), None, 7), (R2, D(2024, 6, 14), False, 6),
               (D(2024, 6, 1), D(2024, 7, 31), True, 3), (R3, D(2025, 9, 14), None, 7)]
    for lo, hi, g, k in windows:
        c = [i for i in idx if free(i) and lo <= quotes[i]["date"] <= hi and (g is None or quotes[i]["grace"] == g)]
        pick_apply(rng, c, lambda i: rep_w_stale[quotes[i]["rep"]], k, f_stale)

    def f_stale_rate(i):
        q = quotes[i]
        alt = price(q, mult_override=880)
        if alt["net"] == q["rule"]["net"]:
            return False
        set_actual(q, alt)
        q["quirk"] = "stale_table"
        q["detail"] = (f"previous contract multiplier 880 per mille used after {CHAIN_RATE_CHANGE}; "
                       f"diff={q['true_net'] - q['rule']['net']}")
        return True
    c = [i for i in idx if free(i) and quotes[i]["cust"]["code"] == "1023"
         and CHAIN_RATE_CHANGE <= quotes[i]["date"] <= D(2024, 12, 15)]
    pick_apply(rng, c, one, 3, f_stale_rate)

    # --- another customer's multiplier copy-pasted
    rep_w_mult = {"田中": 0.6, "渡辺": 1.0, "佐藤": 1.2, "鈴木": 1.4, "高橋": 1.4}

    def f_mult(i):
        q = quotes[i]
        cur = cust_mult(q["cust"], q["date"])
        opts = [m for m in (800, 850, 880, 900, 950, 1000) if m != cur]
        wrong = int(choice(rng, opts))
        alt = price(q, mult_override=wrong)
        if alt["net"] == q["rule"]["net"]:
            return False
        set_actual(q, alt)
        q["quirk"] = "multiplier_copy"
        q["detail"] = f"multiplier {wrong} per mille used instead of {cur}; diff={q['true_net'] - q['rule']['net']}"
        return True
    pick_apply(rng, [i for i in idx if free(i) and quotes[i]["rule"]["product_after_mult"] > 6000],
               lambda i: rep_w_mult[quotes[i]["rep"]], 14, f_mult)

    # --- undocumented manager discount
    rep_w_disc = {"田中": 6.0, "渡辺": 1.0, "佐藤": 1.0, "鈴木": 0.7, "高橋": 0.7}

    def f_disc(i):
        q = quotes[i]
        r = q["rule"]["net"]
        unit = 5000 if r >= 100000 else 1000
        v = int(r * (1 - rng.uniform(0.05, 0.15))) // unit * unit
        if v >= r or v <= 0:
            return False
        set_net(q, v)
        q["quirk"] = "manager_discount"
        q["detail"] = f"undocumented discount {r - v} ({(r - v) / r:.1%}) from rule_net {r}"
        return True
    pick_apply(rng, [i for i in idx if free(i) and quotes[i]["rule"]["net"] >= 15000],
               lambda i: rep_w_disc[quotes[i]["rep"]], 20, f_disc)

    # --- rush surcharge charged, no rush wording
    def f_rush(i):
        q = quotes[i]
        alt = price(q, force_rush=True)
        if alt["net"] == q["rule"]["net"]:
            return False
        set_actual(q, alt)
        q["quirk"] = "rush_unrecorded"
        q["detail"] = f"rush surcharge {alt['rush']} applied without rush wording in remarks"
        return True
    pick_apply(rng, [i for i in idx if free(i) and not quotes[i]["flags"]["rush"]], one, 12, f_rush)

    # --- rounding inconsistency
    def f_round(i):
        q = quotes[i]
        x = q["rule"]["pre_round_total"]
        alts = [("none", x), ("ceil100", cdiv(x, 100) * 100), ("halfup1000", (x + 500) // 1000 * 1000),
                ("floor1000", x // 1000 * 1000), ("halfup100", (x + 50) // 100 * 100), ("floor10", x // 10 * 10)]
        alts = [(n, v) for n, v in alts if v != q["rule"]["net"]]
        if not alts:
            return False
        n, v = alts[int(rng.integers(len(alts)))]
        set_net(q, v)
        q["quirk"] = "rounding_inconsistent"
        q["detail"] = (f"rounded '{n}' instead of '{q['rule']['rounding_rule']}' (pre-round {x}); "
                       f"diff={v - q['rule']['net']}")
        return True
    pick_apply(rng, [i for i in idx if free(i)], one, 15, f_round)

    # --- travel fee forgotten
    def f_travel(i):
        q = quotes[i]
        alt = price(q, omit_travel=True)
        if alt["net"] == q["rule"]["net"]:
            return False
        set_actual(q, alt)
        q["quirk"] = "travel_omitted"
        q["detail"] = f"travel fee {q['rule']['travel']} not charged"
        return True
    pick_apply(rng, [i for i in idx if free(i) and quotes[i]["rule"]["travel"] > 0], one, 10, f_travel)

    # --- tax basis written wrongly
    def f_tax_ex(i):
        q = quotes[i]
        q["written"] = q["true_net"] + q["true_net"] * 10 // 100
        q["quirk"] = "tax_label_error"
        q["detail"] = "tax-inclusive amount written under a tax-exclusive or blank label"
        return True

    def f_tax_in(i):
        q = quotes[i]
        if q["written"] == q["true_net"]:
            return False
        q["written"] = q["true_net"]
        q["quirk"] = "tax_label_error"
        q["detail"] = "pre-tax amount written under the 税込 label"
        return True
    pick_apply(rng, [i for i in idx if free(i) and quotes[i]["basis"] == "excl"], one, 4, f_tax_ex)
    pick_apply(rng, [i for i in idx if free(i) and quotes[i]["basis"] == "incl"], one, 4, f_tax_in)

    # --- digit typo in the written amount
    def f_typo(i):
        q = quotes[i]
        correct = q["written"]
        q["written"], kind = digit_typo(correct, rng)
        q["quirk"] = "digit_typo"
        q["detail"] = f"{kind}: intended {correct}, written {q['written']}"
        return True
    pick_apply(rng, [i for i in idx if free(i)], one, 10, f_typo)
    return quotes


def add_revisions_and_duplicates(quotes, rng):
    extra = []
    kinds = iter(["fix", "spec_change", "negotiated", "spec_change", "negotiated", "fix", "negotiated",
                  "spec_change", "negotiated", "spec_change", "negotiated", "spec_change"])
    pending = {"k": next(kinds)}

    def f_rev(i):
        v1 = quotes[i]
        kind = pending["k"]
        if kind == "fix" and v1["rule"]["travel"] == 0:
            return False
        nd = next_open(v1["date"], int(rng.integers(2, 14)))
        if nd > END or not rep_active(v1["rep"], nd):
            return False
        v2 = dict(v1)
        v2["date"] = nd
        v2["rev_of"] = i
        v2["post_adj"] = 0
        v1["quirk"] = "revision_v1"
        if kind == "spec_change":
            if v1["qty"] > 1 and rng.random() < .5:
                v2["qty"] = max(1, v1["qty"] - int(rng.integers(1, v1["qty"])))
            else:
                v2["qty"] = v1["qty"] + int(rng.integers(1, 4))
            v2["remarks"] = v1["remarks"] + ("、" if v1["remarks"] else "") + choice(
                rng, ["数量変更につき再見積", "数量変更", "再見積（数量変更）"])
            v2["flags"] = parse_flags(v2["remarks"])
            v2["rule"] = price(v2)
            v2["actual"] = v2["rule"]
            finalize_clean(v2)
            v1["detail"] = "superseded by re-issued quote (quantity change); v1 priced by rules"
            v2["quirk"], v2["detail"] = "revision_v2", f"quantity {v1['qty']}->{v2['qty']}; priced by rules"
        elif kind == "negotiated":
            v2["remarks"] = v1["remarks"] + ("、" if v1["remarks"] else "") + "再見積"
            v2["flags"] = parse_flags(v2["remarks"])
            v2["rule"] = price(v2)
            v2["actual"] = v2["rule"]
            r = v2["rule"]["net"]
            v = int(r * rng.uniform(0.90, 0.96)) // 1000 * 1000
            v2["true_net"] = v
            v2["post_adj"] = v - r
            v2["written"] = written_from_net(v, v2["date"], v2["basis"])
            v2["tax_label"] = tax_label_for(v2["rep"], v2["date"], v2["basis"])
            v1["detail"] = "superseded by re-issued quote (negotiated price); v1 priced by rules"
            v2["quirk"], v2["detail"] = "revision_v2", f"negotiated price, undocumented reduction {r - v} from rule_net {r}"
        else:
            v1["actual"] = price(v1, omit_travel=True)
            finalize_clean(v1)
            v2["remarks"] = v1["remarks"] + ("、" if v1["remarks"] else "") + "再見積（訂正）"
            v2["flags"] = parse_flags(v2["remarks"])
            v2["rule"] = price(v2)
            v2["actual"] = v2["rule"]
            finalize_clean(v2)
            v1["detail"] = f"travel fee {v1['rule']['travel']} omitted in v1; corrected in re-issued quote"
            v2["quirk"], v2["detail"] = "revision_v2", "correction of v1 (travel fee added); priced by rules"
        extra.append(v2)
        pending["k"] = next(kinds, None)
        return True
    pick_apply(rng, [i for i, q in enumerate(quotes) if q["quirk"] == "clean"], lambda i: 1.0, 12, f_rev)

    def f_dup(i):
        src = quotes[i]
        dup = dict(src)
        dup["dup_of"] = i
        if rng.random() < 0.4:
            nd = next_open(src["date"], 1)
            if nd <= END and rep_active(src["rep"], nd):
                dup["date"] = nd
        dup["quirk"], dup["detail"] = "duplicate", "same quote entered twice under a new number"
        extra.append(dup)
        return True
    pick_apply(rng, [i for i, q in enumerate(quotes) if q["quirk"] == "clean"], lambda i: 1.0, 8, f_dup)
    return quotes + extra


def assign_numbers(rows, rng):
    for r in rows:
        r["tb"] = float(rng.random())
    rows.sort(key=lambda r: (r["date"], r["tb"]))
    counters = {}
    for r in rows:
        if "rev_of" in r:
            continue
        key = (r["date"].year % 100, r["date"].month)
        counters[key] = counters.get(key, 0) + 1
        r["quote_no"] = f"{key[0]:02d}{key[1]:02d}-{counters[key]:03d}"
    for r in rows:
        if "rev_of" in r:
            r["quote_no"] = r["src_ref"]["quote_no"] + "-2"
    return rows


# ----------------------------------------------------------------------------
# Output: quotes.csv
# ----------------------------------------------------------------------------
CSV_COLS = ["quote_no", "issue_date", "customer_code", "rep", "item", "width_mm", "height_mm", "qty", "unit",
            "thickness_mm", "sides", "media", "laminate", "finish", "lighting", "design", "install",
            "install_height_m", "site", "remarks", "amount", "tax_label"]


def fmt_h(ht):
    return "" if ht is None else f"{ht / 10:.1f}"


def csv_row(r):
    site = site_name(r["site"], r["date"], r["rep"]) if r["site"] is not None else ""
    return [r["quote_no"], r["date"].isoformat(), r["cust"]["code"], r["rep"], ITEM_NAME[r["item"]],
            "" if r["W"] is None else r["W"], r["H"], r["qty"], r["unit"],
            "" if r["thick"] is None else r["thick"], r["sides"], r["media"], r["lam"], r["finish"],
            r["light"], r["design"], "有" if r["install"] else "無",
            fmt_h(r["height_t"]) if r["install"] else "", site, r["remarks"], r["written"], r["tax_label"]]


def write_csv(path, header, rows):
    buf = io.StringIO()
    w = csv.writer(buf, lineterminator="\n")
    w.writerow(header)
    for row in rows:
        w.writerow(row)
    path.write_text(buf.getvalue(), encoding="utf-8")


# ----------------------------------------------------------------------------
# Messy workbooks
# ----------------------------------------------------------------------------
ZEN = str.maketrans("0123456789,.-/()R%:", "０１２３４５６７８９，．－／（）Ｒ％：")
FIXED_TS = {"messy1": "2024-04-12T08:31:00Z", "messy2": "2025-04-07T10:02:00Z", "messy3": "2025-10-03T17:45:00Z"}


def z(s):
    return str(s).translate(ZEN)


def wareki(d, style):
    ry = d.year - 2018
    if style == "kanji":
        return f"令和{ry}年{d.month}月{d.day}日"
    if style == "dot":
        return f"R{ry}.{d.month}.{d.day}"
    if style == "slash":
        return f"R{ry}/{d.month}/{d.day}"
    if style == "zen":
        return z(f"R{ry}.{d.month}.{d.day}")
    if style == "short":
        return f"{ry}/{d.month}/{d.day}"
    raise ValueError(style)


def spec_detail(r):
    it = r["item"]
    if it == "ACM":
        return f"{r['sides']} " + ("ラミなし" if r["lam"] == "なし" else f"{r['lam']}ラミ")
    if it == "IJ":
        return f"{r['media']} " + ("ラミなし" if r["lam"] == "なし" else f"{r['lam']}ラミ")
    if it == "CALP":
        return f"t{r['thick']} {r['finish']}"
    return f"{'LED内照' if r['light'] == 'LED内照' else '非照明'} 両面"


def size_text(r, style=0):
    if r["item"] == "CALP":
        return f"H{r['H']}"
    if style == 1:
        return f"W{r['W']}×H{r['H']}"
    return f"{r['W']}×{r['H']}"


SHORT_ITEM = {"ACM": "アルミ複合板", "CALP": "カルプ文字", "IJ": "IJ出力", "SODE": "袖看板"}


def normalize_xlsx(path, ts):
    with zipfile.ZipFile(path) as zf:
        items = [(i.filename, zf.read(i.filename)) for i in zf.infolist()]
    buf = io.BytesIO()
    y, mo, d_, hh, mi = int(ts[0:4]), int(ts[5:7]), int(ts[8:10]), int(ts[11:13]), int(ts[14:16])
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        for name, data in items:
            if name == "docProps/core.xml":
                data = re.sub(rb"(<dcterms:created[^>]*>)[^<]*(</dcterms:created>)",
                              rb"\g<1>" + ts.encode() + rb"\g<2>", data)
                data = re.sub(rb"(<dcterms:modified[^>]*>)[^<]*(</dcterms:modified>)",
                              rb"\g<1>" + ts.encode() + rb"\g<2>", data)
            zi = zipfile.ZipInfo(name, date_time=(y, mo, d_, hh, mi, 0))
            zi.compress_type = zipfile.ZIP_DEFLATED
            zi.external_attr = 0o644 << 16
            zf.writestr(zi, data)
    path.write_bytes(buf.getvalue())


def new_wb(creator):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.creator = creator
    wb.properties.lastModifiedBy = creator
    return wb


THIN = Side(style="thin")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def messy1(rows, path, rng, mapping):
    """Monthly ledger FY2023 (R5年度), one sheet per month, 和暦 dates, merged headers."""
    wb = new_wb("sato")
    lo, hi = D(2023, 4, 1), D(2024, 3, 31)
    sel = [r for r in rows if lo <= r["date"] <= hi]
    months = []
    for r in sel:
        k = (r["date"].year, r["date"].month)
        if k not in months:
            months.append(k)
    date_style = {"田中": "kanji", "佐藤": "dot", "渡辺": "slash", "鈴木": "cell", "高橋": "zen"}
    no_ref = set(int(i) for i in rng.choice(len(sel), size=3, replace=False))
    for (y, m) in months:
        ws = wb.create_sheet(f"R{y - 2018}年{m}月")
        ws.merge_cells("A1:P1")
        ws["A1"] = z(f"見積台帳　令和{y - 2018}年{m}月分")
        ws["A1"].font = Font(bold=True, size=14)
        ws["A1"].alignment = Alignment(horizontal="center")
        ws["A2"] = "（社内用）"
        for col, title in zip("ABCDEF", ["No.", "見積日", "見積番号", "得意先", "担当", "品名"]):
            ws.merge_cells(f"{col}3:{col}4")
            ws[f"{col}3"] = title
        ws.merge_cells("G3:I3"); ws["G3"] = "仕様"
        ws["G4"], ws["H4"], ws["I4"] = "サイズ", "数量", "詳細"
        ws.merge_cells("J3:L3"); ws["J3"] = "施工"
        ws["J4"], ws["K4"], ws["L4"] = "有無", "高さ", "場所"
        for col, title in zip("MNOP", ["デザイン", "金額", "税", "備考"]):
            ws.merge_cells(f"{col}3:{col}4")
            ws[f"{col}3"] = title
        for row in ws.iter_rows(min_row=3, max_row=4, min_col=1, max_col=16):
            for cell in row:
                cell.alignment = Alignment(horizontal="center", vertical="center")
                cell.font = Font(bold=True)
        for col, wdt in zip("ABCDEFGHIJKLMNOP", [5, 16, 12, 8, 7, 12, 14, 8, 16, 5, 6, 18, 8, 11, 5, 36]):
            ws.column_dimensions[col].width = wdt
        rr = 5
        k = 0
        for gi, r in enumerate(sel):
            if (r["date"].year, r["date"].month) != (y, m):
                continue
            k += 1
            st = date_style[r["rep"]]
            ws.cell(rr, 1, k)
            if st == "cell":
                c = ws.cell(rr, 2, dt.datetime(r["date"].year, r["date"].month, r["date"].day))
                c.number_format = '[$-ja-JP-x-gannen]ggge"年"m"月"d"日";@'
            else:
                ws.cell(rr, 2, wareki(r["date"], st))
            ws.cell(rr, 3, ("No." + r["quote_no"]) if r["rep"] == "鈴木" else r["quote_no"])
            code = r["cust"]["code"]
            ws.cell(rr, 4, "諸口" if code == "9999" else int(code))
            ws.cell(rr, 5, r["rep"])
            ws.cell(rr, 6, SHORT_ITEM[r["item"]])
            stext = size_text(r, 1 if r["rep"] == "田中" else 0)
            ws.cell(rr, 7, z(stext) if r["rep"] == "高橋" else stext)
            ws.cell(rr, 8, f"{r['qty']}{r['unit']}")
            ws.cell(rr, 9, spec_detail(r))
            if r["install"]:
                ws.cell(rr, 10, "○")
                ws.cell(rr, 11, f"{fmt_h(r['height_t'])}m")
            if r["site"] is not None:
                ws.cell(rr, 12, site_name(r["site"], r["date"], r["rep"]))
            ws.cell(rr, 13, r["design"])
            amt = r["written"]
            note = ""
            if gi in no_ref:
                ws.cell(rr, 14, "別紙参照")
                note = "amount cell reads 別紙参照"
            elif r["rep"] == "渡辺":
                ws.cell(rr, 14, f"¥{amt:,}")
            elif r["rep"] == "高橋":
                ws.cell(rr, 14, z(f"{amt:,}"))
            else:
                c = ws.cell(rr, 14, amt)
                c.number_format = "#,##0"
            ws.cell(rr, 15, {"税別": "別", "税抜": "抜", "税込": "込", "": ""}[r["tax_label"]])
            rem = r["remarks"]
            if len(rem) > 22 and r["rep"] in ("佐藤", "渡辺"):
                rem = rem[:22]
                note = (note + "; " if note else "") + "remarks truncated"
            ws.cell(rr, 16, rem)
            mapping.append(dict(file="messy/見積台帳_R5年度.xlsx", sheet=ws.title, rows=str(rr),
                                quote_id=r["quote_no"], kind="ledger_row", note=note))
            r["messy_refs"].append(f"M1:{ws.title}!{rr}")
            rr += 1
        ws.cell(rr, 13, "月計")
        ws.cell(rr, 14, f"=SUM(N5:N{rr - 1})")
        ws.cell(rr, 13).font = Font(bold=True)
    wb.save(path)
    normalize_xlsx(path, FIXED_TS["messy1"])


def line_items(r):
    """Printed-quote style lines derived from the actual (as-quoted) components."""
    a = r["actual"]
    if r["quirk"] == "impossible":
        return [("看板製作一式", spec_detail(r), "1", "式", r["true_net"])], 0, ""
    lines = []
    prod = a["product_final"] + (a["rush"] if not r["flags"]["rush"] else 0)
    lines.append((f"{ITEM_NAME[r['item']]} 製作", f"{size_text(r)} {spec_detail(r)}", str(r["qty"]), r["unit"], prod))
    if r["flags"]["rush"] and a["rush"]:
        lines.append(("特急料金", "", "1", "式", a["rush"]))
    if a["design_fee"]:
        lines.append(("デザイン費", r["design"], "1", "式", a["design_fee"]))
    if r["install"]:
        lines.append(("取付施工費", f"取付高さ{fmt_h(r['height_t'])}m" + ("（夜間）" if r["flags"]["night"] else ""),
                      "1", "式", a["install_base"] + a["night_extra"]))
        if a["height_fee"]:
            lines.append(("高所作業費", "", "1", "式", a["height_fee"]))
        if a["electrical"]:
            lines.append(("電気工事", "", "1", "式", a["electrical"]))
        if a["removal"]:
            lines.append(("既存看板撤去", "", "1", "式", a["removal"]))
        if a["disposal"]:
            lines.append(("産廃処分費", "", "1", "式", a["disposal"]))
    if a["survey"]:
        lines.append(("現地調査費", "", "1", "式", a["survey"]))
    if a["travel"]:
        lines.append(("出張費", site_name(r["site"], r["date"], r["rep"]), "1", "式", a["travel"]))
    s = sum(x[4] for x in lines)
    diff = r["true_net"] - s
    lab = "お値引き" if r["quirk"] in ("manager_discount", "revision_v2") and r.get("post_adj", 0) < 0 else "端数調整"
    return lines, diff, lab


def messy2(rows, path, rng, mapping):
    """Printed quote copies (PDF-like layout), Tanaka & Sato, 2024-10 .. 2025-03."""
    wb = new_wb("tanaka")
    lo, hi = D(2024, 10, 1), D(2025, 3, 31)
    sel = [r for r in rows if lo <= r["date"] <= hi and r["rep"] in ("田中", "佐藤") and r["quirk"] != "duplicate"]
    sheets = {}
    for r in sel:
        key = f"{r['date'].year}.{r['date'].month:02d}"
        if key not in sheets:
            ws = wb.create_sheet(key)
            for col, wdt in zip("ABCDEF", [24, 18, 18, 7, 6, 14]):
                ws.column_dimensions[col].width = wdt
            sheets[key] = [ws, 1]
        ws, rr = sheets[key]
        top = rr
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
        ws.cell(rr, 1, "御　見　積　書").font = Font(bold=True, size=16)
        ws.cell(rr, 1).alignment = Alignment(horizontal="center")
        rr += 1
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
        code = r["cust"]["code"]
        ws.cell(rr, 1, ("諸口　様" if code == "9999" else f"得意先コード {code}　御中"))
        ws.cell(rr, 5, "No.")
        ws.cell(rr, 6, r["quote_no"])
        rr += 1
        ws.cell(rr, 5, "発行日")
        ws.cell(rr, 6, wareki(r["date"], "kanji"))
        rr += 1
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=3)
        ws.cell(rr, 1, f"件名：{ITEM_NAME[r['item']]} 製作" + ("・取付" if r["install"] else ""))
        ws.merge_cells(start_row=rr, start_column=5, end_row=rr, end_column=6)
        ws.cell(rr, 5, f"担当：{r['rep']}")
        rr += 1
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
        ws.cell(rr, 1, "下記の通り御見積申し上げます。")
        rr += 1
        for ci, t in enumerate(["品名", "仕様", "", "数量", "単位", "金額"], start=1):
            ws.cell(rr, ci, t).font = Font(bold=True)
            ws.cell(rr, ci).border = BOX
        ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
        rr += 1
        lines, diff, lab = line_items(r)
        for name, spec, qty, unit, amt in lines:
            ws.cell(rr, 1, name)
            ws.merge_cells(start_row=rr, start_column=2, end_row=rr, end_column=3)
            ws.cell(rr, 2, spec)
            ws.cell(rr, 4, qty)
            ws.cell(rr, 5, unit)
            ws.cell(rr, 6, f"¥{amt:,}").alignment = Alignment(horizontal="right")
            rr += 1
        if diff:
            ws.cell(rr, 1, lab)
            ws.cell(rr, 6, f"-¥{-diff:,}" if diff < 0 else f"¥{diff:,}").alignment = Alignment(horizontal="right")
            rr += 1
        net = r["true_net"]
        tax = net * 10 // 100
        for name, amt in (("小計", net), ("消費税（10%）", tax), ("合計", net + tax)):
            ws.cell(rr, 5, name)
            ws.cell(rr, 6, f"¥{amt:,}").alignment = Alignment(horizontal="right")
            rr += 1
        ws.merge_cells(start_row=rr, start_column=1, end_row=rr, end_column=6)
        ws.cell(rr, 1, f"備考：{r['remarks']}" if r["remarks"] else "備考：")
        rr += 1
        mapping.append(dict(file="messy/見積書控_2024下期.xlsx", sheet=ws.title, rows=f"{top}-{rr - 1}",
                            quote_id=r["quote_no"], kind="printed_block",
                            note="adjustment line present" if diff else ""))
        r["messy_refs"].append(f"M2:{ws.title}!{top}-{rr - 1}")
        rr += 2
        sheets[key][1] = rr
    wb.save(path)
    normalize_xlsx(path, FIXED_TS["messy2"])


def messy3(rows, path, rng, mapping):
    """Per-customer lists for the busiest accounts, 2024-04 .. 2025-09, full-width text, adjustment rows."""
    wb = new_wb("suzuki")
    lo, hi = D(2024, 4, 1), D(2025, 9, 30)
    sel = [r for r in rows if lo <= r["date"] <= hi and r["cust"]["code"] != "9999"]
    counts = {}
    for r in sel:
        counts[r["cust"]["code"]] = counts.get(r["cust"]["code"], 0) + 1
    top = sorted(counts, key=lambda c: (-counts[c], c))[:6]
    top = sorted(top)
    for si, code in enumerate(top):
        style = si % 3
        title = [code, z(code), f"{code}様"][si % 3]
        ws = wb.create_sheet(title)
        numeric = style == 1
        ws.merge_cells("A1:G1")
        ws["A1"] = f"得意先別 見積一覧（{code}）　2024年4月～"
        ws["A1"].font = Font(bold=True, size=13)
        hdr = ["日付", "No", "品名・仕様", "数量", "施工", "金額" if numeric else "金額（円）", "備考"]
        for ci, t in enumerate(hdr, start=1):
            ws.cell(3, ci, t).font = Font(bold=True)
        for col, wdt in zip("ABCDEFG", [12, 12, 40, 8, 26, 14, 40]):
            ws.column_dimensions[col].width = wdt
        rr = 4
        cur_year = None
        first_num_row = rr
        for r in [x for x in sel if x["cust"]["code"] == code]:
            if style == 2 and r["date"].year != cur_year:
                cur_year = r["date"].year
                ws.cell(rr, 1, f"【{cur_year}年】")
                rr += 1
            if style == 0:
                ws.cell(rr, 1, wareki(r["date"], "short"))
            elif style == 1:
                c = ws.cell(rr, 1, dt.datetime(r["date"].year, r["date"].month, r["date"].day))
                c.number_format = "yyyy/m/d"
            else:
                ws.cell(rr, 1, f"{r['date'].month}月{r['date'].day}日")
            ws.cell(rr, 2, r["quote_no"] if numeric else z(r["quote_no"]))
            desc = f"{SHORT_ITEM[r['item']]} {size_text(r)} {spec_detail(r)}"
            ws.cell(rr, 3, desc if numeric else z(desc))
            ws.cell(rr, 4, z(f"{r['qty']}{r['unit']}"))
            if r["install"]:
                ws.cell(rr, 5, z(f"有({fmt_h(r['height_t'])}m) ") + site_name(r["site"], r["date"], r["rep"]))
            else:
                ws.cell(rr, 5, "－")
            a = r["actual"]
            show_adj = (r["basis"] == "excl" and r["quirk"] not in ("impossible", "digit_typo", "tax_label_error")
                        and (r["true_net"] - a["pre_round_total"] <= -300))
            amount = a["pre_round_total"] if show_adj else r["written"]
            if r["quirk"] in ("digit_typo",):
                amount = written_from_net(r["true_net"], r["date"], r["basis"])
            if r["quirk"] == "tax_label_error":
                amount = written_from_net(r["true_net"], r["date"], r["basis"])
            note = ""
            if not show_adj and rng.random() < 0.012:
                bad, kind = digit_typo(amount, rng)
                note = f"transcription error ({kind}): list shows {bad}, quote amount {amount}"
                amount = bad
            if numeric:
                c = ws.cell(rr, 6, amount)
                c.number_format = "#,##0"
            else:
                ws.cell(rr, 6, z(f"{amount:,}") + "円" + ("（税込）" if r["basis"] == "incl" else ""))
            ws.cell(rr, 7, r["remarks"])
            mapping.append(dict(file="messy/得意先別見積一覧.xlsx", sheet=ws.title, rows=str(rr),
                                quote_id=r["quote_no"], kind="list_row", note=note))
            r["messy_refs"].append(f"M3:{ws.title}!{rr}")
            rr += 1
            if show_adj:
                adj = r["true_net"] - a["pre_round_total"]
                ws.cell(rr, 3, "　　調整" if numeric else "　　端数調整・値引")
                if numeric:
                    c = ws.cell(rr, 6, adj)
                    c.number_format = "#,##0"
                else:
                    ws.cell(rr, 6, "－" + z(f"{-adj:,}") + "円")
                mapping.append(dict(file="messy/得意先別見積一覧.xlsx", sheet=ws.title, rows=str(rr),
                                    quote_id=r["quote_no"], kind="adjustment_row", note=""))
                r["messy_refs"].append(f"M3:{ws.title}!{rr}(adj)")
                rr += 1
        if numeric:
            ws.cell(rr, 5, "合計")
            ws.cell(rr, 6, f"=SUM(F{first_num_row}:F{rr - 1})")
    wb.save(path)
    normalize_xlsx(path, FIXED_TS["messy3"])


# ----------------------------------------------------------------------------
# SCHEMA.md (public, neutral)
# ----------------------------------------------------------------------------
SCHEMA_MD = """# quotes.csv — 列定義 / Column definitions

看板製作会社の見積台帳から書き出した見積データです。1行が1枚の見積書に対応します。
Quotes exported from a signage shop's quote ledger. One row corresponds to one issued quote sheet.

- 文字コード / encoding: UTF-8 (no BOM), comma-separated, one header row
- 並び順 / order: issue_date, then quote number as issued
- 行数 / rows: {nrows} (excluding header)
- 空欄 / blank: the field was not written on the quote or does not apply to the item

| column | 項目名 | 内容 | Description |
|---|---|---|---|
| quote_no | 見積番号 | 見積書に記載の番号 | Quote number as written (`YYMM-NNN`; some numbers carry an extra suffix such as `-2`) |
| issue_date | 見積日 | 見積書の発行日 (YYYY-MM-DD) | Issue date of the quote |
| customer_code | 得意先コード | 得意先コード。`9999` は諸口 | Customer account code. `9999` is the shared code for one-off customers (諸口) |
| rep | 担当者 | 営業担当者（姓） | Sales rep who wrote the quote (surname) |
| item | 品目 | アルミ複合板看板 / カルプ文字 / インクジェット出力 / 袖看板 | Product category |
| width_mm | 幅 (mm) | 製作物の幅 | Width of the piece, mm (blank for カルプ文字) |
| height_mm | 高さ (mm) | 製作物の高さ。カルプ文字は文字高 | Height of the piece, mm; for カルプ文字 it is the character height |
| qty | 数量 | 数量 | Quantity |
| unit | 単位 | 枚 / 文字 / 台 | Unit of the quantity (sheets / characters / units) |
| thickness_mm | 厚み (mm) | 板厚または文字の厚み | Board thickness (アルミ複合板) or letter thickness (カルプ文字), mm |
| sides | 面 | 片面 / 両面 | Single-sided / double-sided |
| media | メディア | 出力メディア（インクジェット出力のみ） | Print media (インクジェット出力 only) |
| laminate | ラミネート | グロス / マット / なし | Laminate finish |
| finish | 仕上げ | シート貼り / 塗装（カルプ文字のみ） | Face finish (カルプ文字 only) |
| lighting | 照明 | LED内照 / なし（袖看板のみ） | Illumination (袖看板 only) |
| design | デザイン | 支給 / 修正 / 新規 | Artwork: customer-supplied / adjustment of supplied artwork / new design |
| install | 施工 | 有 / 無 | Whether on-site installation is part of the quote |
| install_height_m | 取付高さ (m) | 取付位置の高さ | Mounting height in metres (blank when not installed) |
| site | 施工場所 | 現場の所在地（記載どおり） | Site location as written on the quote |
| remarks | 備考 | 備考欄の記載（原文のまま） | Free-text remarks, verbatim |
| amount | 金額 (円) | 見積書に記載の金額 | Quoted amount in JPY, exactly as written on the quote |
| tax_label | 税表示 | 金額欄の税表示（記載どおり） | Tax wording printed next to the amount (税別 / 税抜 / 税込 / blank), as written |

## messy/

同じ見積の一部を、社内で別途作成されていた台帳・控え・一覧の元の形式のまま収録しています。
Three workbooks kept in their original in-house layouts. Each holds an alternative record of a subset of
the quotes above; a quote may appear in none, one or several of them.

| file | 内容 | Contents |
|---|---|---|
| messy/見積台帳_R5年度.xlsx | 月別見積台帳（令和5年度） | Monthly quote ledger, fiscal year R5 (one sheet per month) |
| messy/見積書控_2024下期.xlsx | 発行した見積書の控え | File copies of printed quote sheets |
| messy/得意先別見積一覧.xlsx | 得意先別の見積一覧 | Per-customer quote lists |
"""


# ----------------------------------------------------------------------------
# Truth
# ----------------------------------------------------------------------------
TRUTH_COLS = (["quote_id", "issue_date", "customer_code", "rep", "item", "quirk", "quirk_detail",
               "rule_net", "true_net", "written_amount", "written_tax_label", "tax_basis", "flags_from_remarks"]
              + ["rule_" + k for k in ["table_version"] + COMP_KEYS + ["rounding_rule", "rounding_adj"]]
              + ["actual_table_version", "actual_mult", "actual_rush", "actual_travel", "actual_pre_round_total",
                 "post_rounding_adjustment", "related_quote", "messy_refs"])


def truth_row(r):
    ru, a = r["rule"], r["actual"]
    rel = ""
    if "rev_of" in r:
        rel = "v1=" + r["src_ref"]["quote_no"]
    elif r["quirk"] == "revision_v1":
        rel = "v2=" + r["quote_no"] + "-2"
    elif "dup_of" in r:
        rel = "original=" + r["src_ref"]["quote_no"]
    flags = ";".join(k for k in FLAG_KEYS if r["flags"][k])
    row = [r["quote_no"], r["date"].isoformat(), r["cust"]["code"], r["rep"], r["item"], r["quirk"], r["detail"],
           ru["net"], r["true_net"], r["written"], r["tax_label"], r["basis"], flags]
    row += [ru["table_version"]] + [ru[k] for k in COMP_KEYS] + [ru["rounding_rule"], ru["rounding_adj"]]
    row += [a["table_version"], a["mult"], a["rush"], a["travel"], a["pre_round_total"],
            r["true_net"] - a["pre_round_total"], rel, " | ".join(r["messy_refs"])]
    return row


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def sha256(p):
    return hashlib.sha256(Path(p).read_bytes()).hexdigest()


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--repo-out", default=str(DEFAULT_REPO_OUT))
    ap.add_argument("--sealed-out", default=str(DEFAULT_SEALED_OUT))
    args = ap.parse_args()
    repo, sealed = Path(args.repo_out), Path(args.sealed_out)
    (repo / "messy").mkdir(parents=True, exist_ok=True)
    sealed.mkdir(parents=True, exist_ok=True)

    for t in NOISE_GENERAL + NOISE_INSTALL + DECOY_RUSH + DECOY_SURVEY + DECOY_REMOVAL:
        assert not any(parse_flags(t)[k] for k in FLAG_KEYS if k not in ()), t

    custs = build_customers(RNG_CUST)
    base = generate_base(custs, RNG_Q)
    base = apply_quirks(base, RNG_QUIRK)
    rows = add_revisions_and_duplicates(base, RNG_QUIRK)
    for r in rows:
        if "rev_of" in r:
            r["src_ref"] = rows[r["rev_of"]]
        if "dup_of" in r:
            r["src_ref"] = rows[r["dup_of"]]
        r["messy_refs"] = []
    rows = assign_numbers(rows, RNG_ORDER)

    write_csv(repo / "quotes.csv", CSV_COLS, [csv_row(r) for r in rows])
    (repo / "SCHEMA.md").write_text(SCHEMA_MD.format(nrows=len(rows)), encoding="utf-8")

    mapping = []
    messy1(rows, repo / "messy" / "見積台帳_R5年度.xlsx", RNG_MESSY, mapping)
    messy2(rows, repo / "messy" / "見積書控_2024下期.xlsx", RNG_MESSY, mapping)
    messy3(rows, repo / "messy" / "得意先別見積一覧.xlsx", RNG_MESSY, mapping)

    truth_sorted = sorted(rows, key=lambda r: r["quote_no"])
    write_csv(sealed / "truth.csv", TRUTH_COLS, [truth_row(r) for r in truth_sorted])
    write_csv(sealed / "messy_map.csv", ["file", "sheet", "rows", "quote_id", "kind", "note"],
              [[m["file"], m["sheet"], m["rows"], m["quote_id"], m["kind"], m["note"]] for m in mapping])

    lines = [f"# quotes.csv data rows (excluding header): {len(rows)}"]
    for name, p in (("generator.py", HERE / "generator.py"), ("RULES.md", HERE / "RULES.md"),
                    ("truth.csv", sealed / "truth.csv"), ("messy_map.csv", sealed / "messy_map.csv")):
        if p.exists():
            lines.append(f"{sha256(p)}  {name}")
    (repo / "SEALED.sha256").write_text("\n".join(lines) + "\n", encoding="utf-8")

    from collections import Counter
    cnt = Counter(r["quirk"] for r in rows)
    print("rows", len(rows), dict(sorted(cnt.items())))


if __name__ == "__main__":
    main()
