import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).parent
sys.path.insert(0, str(HERE.parent))

import scraper  # noqa: E402

PAGE_URL = "https://books.toscrape.com/catalogue/page-1.html"


class ParseListPageTest(unittest.TestCase):
    def test_first_page(self):
        html = (HERE / "fixtures/page-1.html").read_text(encoding="utf-8")
        books, next_url = scraper.parse_list_page(html, PAGE_URL)
        self.assertEqual(len(books), 20)
        first = books[0]
        self.assertEqual(first.title, "A Light in the Attic")
        self.assertAlmostEqual(first.price_gbp, 51.77)
        self.assertEqual(first.rating, 3)
        self.assertTrue(first.in_stock)
        self.assertEqual(first.url, "https://books.toscrape.com/catalogue/a-light-in-the-attic_1000/index.html")
        self.assertEqual(next_url, "https://books.toscrape.com/catalogue/page-2.html")

    def test_last_page_has_no_next(self):
        html = (HERE / "fixtures/page-50.html").read_text(encoding="utf-8")
        books, next_url = scraper.parse_list_page(html, PAGE_URL)
        self.assertEqual(len(books), 20)
        self.assertIsNone(next_url)


class WriteExcelTest(unittest.TestCase):
    def test_sheets_and_summary(self):
        books = [
            scraper.Book("A", 10.0, 5, True, "u1"),
            scraper.Book("B", 20.0, 5, False, "u2"),
            scraper.Book("C", 30.0, 1, True, "u3"),
        ]
        with tempfile.TemporaryDirectory() as d:
            path = f"{d}/out.xlsx"
            scraper.write_excel(books, path)
            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames, ["一覧", "評価別集計"])
            self.assertEqual(wb["一覧"].max_row, 4)
            rows = list(wb["評価別集計"].iter_rows(min_row=2, values_only=True))
            self.assertEqual(rows, [(5, 2, 15.0, 10.0, 20.0), (1, 1, 30.0, 30.0, 30.0)])


if __name__ == "__main__":
    unittest.main()
