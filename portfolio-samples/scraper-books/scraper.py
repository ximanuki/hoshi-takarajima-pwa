"""書籍一覧サイトから商品情報を集めて Excel にまとめるツール。

対象は練習用に公開されている https://books.toscrape.com/ 。
実案件ではここを依頼主のサイトに差し替え、parse_list_page() のセレクタだけ直せば動く構成にしている。

使い方:
    python scraper.py --pages 3 --out books.xlsx
"""

from __future__ import annotations

import argparse
import sys
import time
from dataclasses import dataclass, asdict
from urllib.parse import urljoin

import requests
from bs4 import BeautifulSoup
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

BASE_URL = "https://books.toscrape.com/catalogue/"
USER_AGENT = "portfolio-sample-scraper/1.0 (+contact: your-email@example.com)"
RATINGS = {"One": 1, "Two": 2, "Three": 3, "Four": 4, "Five": 5}


@dataclass
class Book:
    title: str
    price_gbp: float
    rating: int
    in_stock: bool
    url: str


def parse_list_page(html: str, page_url: str) -> tuple[list[Book], str | None]:
    """一覧ページ1枚から商品と「次へ」のURLを取り出す。"""
    soup = BeautifulSoup(html, "html.parser")
    books: list[Book] = []
    for pod in soup.select("article.product_pod"):
        link = pod.select_one("h3 a")
        price_text = pod.select_one("p.price_color").get_text(strip=True)
        rating_class = next(
            (c for c in pod.select_one("p.star-rating")["class"] if c in RATINGS), None
        )
        books.append(
            Book(
                title=link["title"],
                price_gbp=float(price_text.lstrip("Â£")),
                rating=RATINGS.get(rating_class, 0),
                in_stock="In stock" in pod.select_one("p.availability").get_text(),
                url=urljoin(page_url, link["href"]),
            )
        )
    next_link = soup.select_one("li.next a")
    next_url = urljoin(page_url, next_link["href"]) if next_link else None
    return books, next_url


def crawl(max_pages: int, delay: float, session: requests.Session | None = None) -> list[Book]:
    """一覧ページを順にたどる。相手サーバーに負担をかけないよう、1リクエストごとに待つ。"""
    session = session or requests.Session()
    session.headers["User-Agent"] = USER_AGENT
    url: str | None = urljoin(BASE_URL, "page-1.html")
    books: list[Book] = []
    for page_no in range(1, max_pages + 1):
        if url is None:
            break
        for attempt in range(3):
            try:
                res = session.get(url, timeout=15)
                res.raise_for_status()
                break
            except requests.RequestException as e:
                if attempt == 2:
                    raise
                print(f"  再試行 {attempt + 1}/2: {e}", file=sys.stderr)
                time.sleep(delay * 2)
        res.encoding = "utf-8"
        page_books, url = parse_list_page(res.text, res.url)
        books.extend(page_books)
        print(f"{page_no}ページ目: {len(page_books)}件（累計 {len(books)}件）")
        if url:
            time.sleep(delay)
    return books


def write_excel(books: list[Book], path: str) -> None:
    """一覧シートと評価別の集計シートを作る。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "一覧"
    headers = ["タイトル", "価格(£)", "評価(★)", "在庫", "URL"]
    ws.append(headers)
    for b in books:
        ws.append([b.title, b.price_gbp, b.rating, "あり" if b.in_stock else "なし", b.url])
    style_header(ws, len(headers))
    for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
        row[0].number_format = "#,##0.00"
    for col, width in zip("ABCDE", (50, 10, 9, 7, 60)):
        ws.column_dimensions[col].width = width
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    summary = wb.create_sheet("評価別集計")
    summary.append(["評価(★)", "件数", "平均価格(£)", "最安(£)", "最高(£)"])
    for stars in range(5, 0, -1):
        prices = [b.price_gbp for b in books if b.rating == stars]
        if prices:
            summary.append([stars, len(prices), round(sum(prices) / len(prices), 2), min(prices), max(prices)])
    style_header(summary, 5)
    for col in range(1, 6):
        summary.column_dimensions[get_column_letter(col)].width = 13
    wb.save(path)


def style_header(ws, ncols: int) -> None:
    fill = PatternFill("solid", fgColor="3D4FD6")
    for col in range(1, ncols + 1):
        cell = ws.cell(row=1, column=col)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")


def main(argv: list[str] | None = None) -> int:
    p = argparse.ArgumentParser(description="書籍一覧を収集して Excel に出力します")
    p.add_argument("--pages", type=int, default=3, help="たどる一覧ページ数（最大50）")
    p.add_argument("--delay", type=float, default=1.0, help="リクエスト間の待ち秒数")
    p.add_argument("--out", default="books.xlsx", help="出力する Excel ファイル")
    args = p.parse_args(argv)

    books = crawl(min(args.pages, 50), args.delay)
    write_excel(books, args.out)
    print(f"{len(books)}件を {args.out} に保存しました")
    return 0


if __name__ == "__main__":
    sys.exit(main())
